"""Выгрузка исходных данных расписания в Excel для правки завучами.

  python export_data.py

Файл не содержит расписания — только данные, по которым оно строится: нагрузка, классы,
учителя, совместители, класс-комплекты, кабинеты, ДО/ЛУО, правила и ручные правки.
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dataset import DAY_LABELS, DAYS, PERIODS, build, grade_of, load
from export_xlsx import room_label, short_name
from import_db import COMBINED
from run_variants import XLSX_DIR

THIN = Side(style="thin", color="AAAAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DDE4EE")
NOTE = PatternFill("solid", fgColor="FFF8E1")
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")


def sheet(wb, title: str, hint: str, head: list[str], rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = title
    c = ws.cell(1, 1, hint)
    c.alignment = WRAP
    c.fill = NOTE
    c.font = Font(italic=True, size=10)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(head)))
    ws.row_dimensions[1].height = 15 * (1 + len(hint) // 110)
    for j, h in enumerate(head):
        cell = ws.cell(2, 1 + j, h)
        cell.fill = HEAD
        cell.font = Font(bold=True, size=10)
        cell.alignment = WRAP_C
        cell.border = BORDER
    ws.row_dimensions[2].height = 32
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = ws.cell(3 + i, 1 + j, v)
            cell.alignment = WRAP
            cell.border = BORDER
            cell.font = Font(size=10)
    for j, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(1 + j)].width = w
    ws.freeze_panes = ws.cell(3, 1)
    return ws


def availability_text(av) -> str:
    if all(all(r) for r in av):
        return "всю неделю"
    parts = []
    for d in range(DAYS):
        ps = [p + 1 for p in range(PERIODS) if av[d][p]]
        if not ps:
            continue
        if ps == list(range(1, PERIONS := PERIODS + 1)):
            parts.append(f"{DAY_LABELS[d]}: весь день")
        else:
            rng, start = [], ps[0]
            for a, b in zip(ps, ps[1:] + [None]):
                if b != a + 1:
                    rng.append(f"{start}–{a}" if start != a else f"{a}")
                    start = b
            parts.append(f"{DAY_LABELS[d]}: {', '.join(rng)} ур.")
    return "; ".join(parts) or "нет"


def main():
    s = build()
    classes_raw = {c["class_name"]: c for c in load("classes.json")["classes"] if c["status"] == "active"}
    room_rules = load("room_rules.json")
    joint = load("joint_lessons.json")["joint_lessons"]
    subgroup_rules = load("subgroup_rules.json")
    overrides = load("overrides.json")
    aisana = load("special_ed_schedule.json")["students"][0]
    do_sched = load("do_scheduling.json")
    luo = load("luo_inclusion.json")
    weights = load("subject_weights.json")
    fixed_types = room_rules["fixed_room_types"]
    base_room = room_rules["teacher_base_room"]

    joint_by_cs = defaultdict(list)
    for j in joint:
        subj_by_class = j.get("subjects") or {c: j["subject"] for c in j["classes"]}
        for c in j["classes"]:
            joint_by_cs[(c, subj_by_class[c])].append(j)

    wb = Workbook()

    # ---------- 1. Нагрузка ----------
    rows = []
    for e in sorted(s.curriculum_entries, key=lambda e: (grade_of(e["class_name"]), e["class_name"],
                                                         e["subject_name"])):
        c, sj = e["class_name"], e["subject_name"]
        tt = e["teachers"]
        teachers = "; ".join(f"{t['fio']}" + (f" ({t['subgroup']})" if t.get("subgroup") else "") for t in tt)
        js = joint_by_cs.get((c, sj), [])
        joint_txt = "; ".join(f"{'+'.join(x['classes'])} — {x['hours_per_week']} ч" for x in js)
        room = fixed_types.get(sj, "")
        room = room_label(s, room) if room.startswith("room_") else {"Gym": "спортзал / малый зал",
                                                                     "Workshop": "мастерские №29 / №39",
                                                                     "Informatics": "каб. информатики (№26)",
                                                                     "ChemistryLab": "каб. химии (№17)",
                                                                     "PhysicsLab": "каб. физики (№16)",
                                                                     "BiologyLab": "каб. биологии (№19)"}.get(room, "")
        manual = "да" if e.get("manual") else ""
        rows.append([c, sj, e["weekly_hours"], teachers,
                     "да" if len({t.get("subgroup") for t in tt if t.get("subgroup")}) > 1 else "",
                     joint_txt, room, manual])
    sheet(wb, "Нагрузка",
          "Главный лист: что и сколько ставим в расписание. Правьте часы, учителей, подгруппы прямо в ячейках. "
          "Чтобы убрать предмет — удалите строку; чтобы добавить — допишите строку снизу. "
          "Пример: «9Б | Художественный труд | 1 | Романишина И.М. (Девочки); Жантимиров А.К. (Мальчики)» — "
          "1 час в неделю, класс делится на две группы, у каждой свой учитель.",
          ["Класс", "Предмет", "Часов в неделю", "Учитель (и подгруппа)", "Деление на подгруппы",
           "Класс-комплект (с кем вместе)", "Кабинет обязателен", "Добавлено вручную"],
          rows, [8, 34, 9, 40, 12, 26, 24, 12])

    # ---------- 2. Классы ----------
    rows = []
    for c in s.classes:
        raw = classes_raw.get(c, {})
        hours = sum(e["weekly_hours"] for e in s.curriculum_entries if e["class_name"] == c)
        home = room_rules["home_rooms"].get(c)
        rows.append([c, "казахский" if s.class_lang.get(c) == "kz" else "русский", raw.get("headcount", ""),
                     raw.get("class_teacher", ""), hours, s.class_max_lessons[c],
                     "ЛУО" if s.class_type.get(c) == "luo" else "",
                     room_label(s, home) if home else ""])
    sheet(wb, "Классы",
          "Список классов. «Предел уроков в день» — сколько уроков максимум можно поставить (1–4 классы — 5, "
          "5–9 — 7, 10–11 — 8; у 3А и 4Б по 6, потому что их часы не помещаются в 5 уроков; у 10А — 7 по "
          "решению школы). Правьте численность, классного руководителя, предел уроков.",
          ["Класс", "Язык обучения", "Учеников", "Классный руководитель", "Часов в неделю",
           "Предел уроков в день", "Особый статус", "Свой кабинет"],
          rows, [8, 14, 10, 32, 12, 12, 12, 26])

    # ---------- 3. Учителя ----------
    by_teacher = defaultdict(lambda: defaultdict(list))
    hours_t = defaultdict(int)
    for e in s.curriculum_entries:
        for t in e["teachers"]:
            by_teacher[t["fio"]][e["subject_name"]].append(e["class_name"])
            hours_t[t["fio"]] += e["weekly_hours"]
    rows = []
    for t in sorted(s.teachers):
        subj_txt = "; ".join(f"{sj} — {', '.join(sorted(set(cl), key=lambda c: (grade_of(c), c)))}"
                             for sj, cl in sorted(by_teacher[t].items()))
        own = [r for r in room_rules["teacher_rooms"].get(t, []) if r in s.rooms]
        rows.append([t, subj_txt, hours_t[t], "да" if t in COMBINED else "",
                     availability_text(s.availability[t]),
                     ", ".join(room_label(s, r) for r in own) if own else "нет (ходит по свободным)"])
    sheet(wb, "Учителя",
          "Нагрузка каждого учителя. В колонке «Предметы и классы» записано в виде «Алгебра — 8А, 8Б, 9Б; "
          "Геометрия — 7А, 10А». «Когда может работать» правьте, если у учителя есть методдни или он приезжает "
          "не каждый день: например «Вт: весь день; Пт: 1–5 ур.».",
          ["Учитель", "Предметы и классы", "Часов в неделю", "Совместитель",
           "Когда может работать", "Закреплённый кабинет"],
          rows, [30, 70, 10, 12, 34, 26])

    # ---------- 4. Совместители ----------
    rows = []
    for t in sorted(COMBINED | {n for n in s.teachers if not all(all(r) for r in s.availability[n])}):
        if t not in s.teachers:
            continue
        subj_txt = "; ".join(f"{sj} — {', '.join(sorted(set(cl), key=lambda c: (grade_of(c), c)))}"
                             for sj, cl in sorted(by_teacher[t].items()))
        note = []
        if t in ("Балташева Зару Такановна", "Аятова Анаргуль Жусуповна"):
            note.append("2 урока у одного класса в день ставятся подряд; профильный кабинет — если свободен")
        if t == "Жантимиров Эльдар Асхатович":
            note.append("окна допускаются: уроки во 2А и 4А стоят в разных концах вторника")
        elif t in COMBINED:
            note.append("окон быть не должно — уроки идут подряд")
        rows.append([t, "да" if t in COMBINED else "нет (есть ограничения по времени)",
                     availability_text(s.availability[t]), hours_t[t], subj_txt, "; ".join(note)])
    sheet(wb, "Совместители",
          "Учителя, которые бывают не всю неделю. Это самое узкое место расписания: чем шире их доступность, "
          "тем меньше окон у остальных. Правьте колонку «Когда может работать» — формат «Вт: 1–8 ур.; Пт: 3–6 ур.».",
          ["Учитель", "Совместитель", "Когда может работать", "Часов в неделю", "Предметы и классы", "Особенности"],
          rows, [30, 14, 34, 10, 56, 40])

    # ---------- 5. Класс-комплекты и деления ----------
    rows = []
    for j in joint:
        subj_by_class = j.get("subjects") or {c: j["subject"] for c in j["classes"]}
        subj = "; ".join(f"{c}: {sj}" for c, sj in subj_by_class.items()) if j.get("subjects") else j["subject"]
        teachers = set()
        for c, sj in subj_by_class.items():
            for e in s.curriculum_entries:
                if e["class_name"] == c and e["subject_name"] == sj:
                    teachers |= {t["fio"] for t in e["teachers"]}
        rows.append(["класс-комплект", " + ".join(j["classes"]), subj, j["hours_per_week"],
                     "; ".join(sorted(teachers)), j.get("comment", "")])
    for e in s.curriculum_entries:
        subs = [t for t in e["teachers"] if t.get("subgroup")]
        if len(subs) > 1:
            rows.append(["деление", e["class_name"], e["subject_name"], e["weekly_hours"],
                         "; ".join(f"{t['subgroup']} — {t['fio']}" for t in subs), ""])
    for cp in subgroup_rules["cross_pairs"]:
        rows.append(["перекрёстно", " и ".join(cp["classes"]), "информатика ↔ английский",
                     sum(x["count_per_week"] for x in cp["combos"]), "Бедрина М.Н.; Алекпарова А.Ж.; Акауова И.Н.",
                     cp.get("comment", "")])
    sheet(wb, "Комплекты и деления",
          "Кто с кем учится вместе и где класс делится на группы. Класс-комплект — уроки идут одновременно у "
          "нескольких классов у одного учителя. Деление — класс делится на две группы, у каждой свой урок в "
          "одно и то же время. Перекрёстно (6Б, 9Б) — у одной подгруппы информатика, у другой английский.",
          ["Тип", "Классы", "Предмет(ы)", "Часов", "Учителя", "Примечание"], rows, [16, 16, 44, 8, 44, 40])

    # ---------- 6. Кабинеты ----------
    rows = []
    inv = defaultdict(list)
    for sj, rt in fixed_types.items():
        inv[rt].append(sj)
    for rid, r in s.rooms.items():
        who = [t for t, rs in room_rules["teacher_rooms"].items() if rid in rs]
        subj = inv.get(r["room_type"], []) + inv.get(rid, [])
        rows.append([room_label(s, rid), r["name"], r["room_type"], r.get("capacity", ""),
                     "; ".join(who), "; ".join(sorted(set(subj)))])
    sheet(wb, "Кабинеты",
          "Кабинеты школы. «Закреплён за» — учитель ведёт здесь свои уроки и не переходит. «Обязателен для» — "
          "предметы, которые можно вести только здесь. Физкультура: большие классы — спортзал, маленькие "
          "(до 8 детей на уроке) — малый зал. Музыка — актовый зал.",
          ["Кабинет", "Полное название", "Тип", "Мест", "Закреплён за", "Обязателен для предметов"],
          rows, [22, 34, 14, 7, 30, 40])

    # ---------- 7. ДО и ЛУО ----------
    rows = []
    for dn, items in aisana["schedule"].items():
        for it in items:
            rows.append([aisana["student"], aisana["class"], "ДО (на дому)", dn, it["period"], it["subject"],
                         it["teacher"], "фиксировано, менять нельзя без согласования"])
    for st in do_sched.get("excluded_for_now", []):
        rows.append([st["student"], st["class"], "ДО / ООП", "", "", "", "", f"в расписание не входит: {st['reason']}"])
    for st in luo["students"]:
        rows.append([st["student"], st["host_class"], "ЛУО (инклюзия)", "", "", "", "",
                     "ходит на уроки класса, отдельных уроков нет"])
    for w in luo["whole_class_luo"]:
        rows.append(["весь класс", w["class"], "ЛУО-класс", "", "", "", "", w["note"]])
    sheet(wb, "ДО и ЛУО",
          "Надомное обучение и ЛУО. В расписание входит только Тургумбаева Айсана: её уроки закреплены и "
          "занимают учителей в эти часы. Остальные ученики ДО и ООП в расписание не включены.",
          ["Ученик", "Класс", "Форма", "День", "Урок", "Предмет", "Учитель", "Примечание"],
          rows, [26, 8, 16, 7, 7, 30, 30, 44])

    # ---------- 8. Правила ----------
    li, ld = weights["language_independent"], weights["language_dependent"]
    hard = [
        ("Окна у детей", "Запрещены: уроки класса идут подряд, без пустых уроков между ними."),
        ("Начало дня", "Все классы начинают с 1-го урока. Исключение — 11А."),
        ("Предел уроков в день", "1–4 классы — 5 (3А и 4Б — 6), 5–9 — 7, 10–11 — 8, 10А — 7. См. лист «Классы»."),
        ("Один предмет в день", "Не больше одного урока предмета в день. Исключения: обучение грамоте в 1А и 1Б "
                                "(6 ч в неделю) и предметы совместителей, которые бывают один день — у них "
                                "2 урока подряд."),
        ("Сдвоенные уроки", "Запрещены, кроме случая выше."),
        ("Математика", "Не позже 5-го урока, в понедельник и пятницу — не позже 4-го. В понедельник ставится "
                       "алгебра, во вторник геометрия."),
        ("Языки в 1–4 классах", "Не больше двух разных языков (русский, казахский, английский) в один день."),
        ("Совместители", "Уроки идут подряд, без окон (кроме Жантимирова Э.А., см. лист «Совместители»)."),
        ("Пащенко Н.А.", "Окон нет."),
        ("Кабинеты", "Учитель со своим кабинетом ведёт у себя. Химия — только в кабинете химии. Физкультура — "
                     "спортзал или малый зал. Информатика — каб. 26. Труды — мастерские. Музыка — актовый зал."),
        ("ЛУО", "6В — не позже 6-го урока."),
        ("Равномерность", "У 1–4 и 11-х классов число уроков по дням отличается не больше чем на 1."),
    ]
    soft = [
        ("Кривая нагрузки (СанПиН)", "5–11 классы: Пн 7, Вт 11, Ср 11, Чт 9, Пт 7. 1–4 классы: Пн 8, Вт 11, "
                                     "Ср 6, Чт 11, Пт 8 — среда облегчённая."),
        ("Трудные предметы", "Ставятся 2–4 уроком; на 1-м, 5-м и последних — со штрафом."),
        ("Чередование", "Алгебра и геометрия, язык и литература, история Казахстана и всемирная история не "
                        "должны стоять в один день и два дня подряд."),
        ("Окна учителей", "Минимизируются."),
        ("8-й урок", "Избегается."),
        ("Физкультура", "Большие классы — спортзал, маленькие — малый зал."),
    ]
    rows = [["ЖЁСТКОЕ", k, v] for k, v in hard] + [["ЖЕЛАТЕЛЬНОЕ", k, v] for k, v in soft]
    rows.append(["СПРАВОЧНО", "Баллы трудности предметов (СанПиН)",
                 "; ".join(f"{k} — {v}" for k, v in sorted(li.items(), key=lambda kv: -kv[1])[:12]) +
                 ". Языки зависят от языка обучения: " +
                 "; ".join(f"{k}: каз. кл. {v['kz']}, рус. кл. {v['ru']}" for k, v in list(ld.items())[:4])])
    sheet(wb, "Правила",
          "Условия, по которым строится расписание. ЖЁСТКОЕ — нарушать нельзя, расписание с нарушением не "
          "принимается. ЖЕЛАТЕЛЬНОЕ — соблюдается по возможности. Если какое-то правило школе не нужно или "
          "нужно другое — напишите об этом прямо в таблице, в свободной строке снизу.",
          ["Тип", "Правило", "Как применяется"], rows, [16, 34, 100])

    # ---------- 9. Отличия от БилимКласса ----------
    rows = []
    for h in overrides.get("hours_overrides", []):
        rows.append(["часы", h["class"], h["subject"], f"{h['weekly_hours']} ч", h.get("note", "")])
    for a in overrides.get("assignments", []):
        if a.get("note"):
            rows.append(["учитель", a["class"], a["subject"],
                         "; ".join(t if isinstance(t, str) else f"{t['fio']} ({t['subgroup']})" for t in a["teachers"]),
                         a["note"]])
    for x in overrides.get("extra_entries", []):
        rows.append(["новый предмет", x["class"], x["subject"], f"{x['weekly_hours']} ч", x.get("note", "")])
    rows.append(["класс-комплект", "5А, 6А, 8Ә", "Физическая культура", "3 ч вместе", "по указанию завуча"])
    rows.append(["класс-комплект", "2А, 4А", "Труд, физкультура, ИЗО, музыка", "вместе", "по указанию завуча"])
    rows.append(["кабинет", "—", "Музыка", "актовый зал", "по указанию завуча"])
    rows.append(["кабинет", "—", "Кабинет робототехники (№15)", "закреплён за Бучковой И.О.", "по указанию завуча"])
    rows.append(["доступность", "—", "Пащенко Н.А., Ажикенова М.А.",
                 "закрыты только их реальные уроки с Айсаной", "в файле доступности был закрыт весь блок 2–3 уроков"])
    sheet(wb, "Отличия от БилимКласса",
          "Ручные правки: чем наши данные отличаются от выгрузки БилимКласса. Если правка неверна — исправьте "
          "или удалите строку, и напишите, как правильно.",
          ["Что изменено", "Класс", "Предмет / кого касается", "Стало", "Причина"], rows, [18, 14, 40, 34, 60])

    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    path = XLSX_DIR / "Данные_для_расписания.xlsx"
    wb.save(path)
    print("записан", path)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
