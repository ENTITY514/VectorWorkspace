"""Печатная сетка расписания — в стиле школьного бланка (Песчанская СОШ №1), без цветов.

Классы (или ученики ДО/ЛУО) — колонки, дни — блоки строк слева с уроком, в ячейке только
предмет (без учителя и кабинета) — как официальный бланк с шапкой согласования («Согласовано»
председатель родкомитета Чоп Н.А., «Утверждаю» директор Мельников П.В., составила Квашнина Е.В.).

    py -3.13 export_print.py V33
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from dataset import DAY_LABELS, DAYS, PERIODS, build, grade_of, load
from export_xlsx import _to_min
from run_variants import OUT, XLSX_DIR

THIN = Side(style="thin", color="000000")
MED = Side(style="medium", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_DAY_TOP = Border(left=THIN, right=THIN, top=MED, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
BOLD = Font(bold=True)

FONT_AGREE_TITLE = Font(name="Calibri", bold=True, size=10)
FONT_AGREE_BODY = Font(name="Calibri", size=9)
FONT_AGREE_NAME = Font(name="Calibri", bold=True, size=10)
FONT_TITLE_MAIN = Font(name="Calibri", bold=True, size=13)
FONT_TITLE_SUB = Font(name="Calibri", bold=True, size=10)
FONT_FOOTER = Font(name="Calibri", bold=True, size=10)

HEADER_INFO = {
    "agree": {
        "title": "«Согласовано»",
        "role": "Председатель родительского комитета:",
        "name": "Чоп Н.А.",
    },
    "approve": {
        "title": "«Утверждаю»",
        "role": "Директор школы:",
        "name": "Мельников П.В.",
    },
    "title": "Расписание уроков на I полугодие",
    "subtitle": "КГУ «Песчанская СОШ №1» 2026-2027 учебный год",
    "compiler": "Составила: ЗДУВР Квашнина Е.В.",
}

# Официальные сокращения предметов для бланка расписания (по образцу школы)
PRINT_SHORT_SUBJECT = {
    "Казахский язык и литература": "КиЛ",
    "История Казахстана": "История Каз.",
    "Всемирная история": "Всем.история",
    "Русская литература": "Русская лит.",
    "Русский язык и литература": "Рус. язык и лит.",
    "Алгебра и начала анализа": "Алгебра и начало ан.",
    "Основы предпринимательства и бизнеса": "ОПиБ",
    "Начальная военная подготовка": "НВТиП",
    "Иностранный язык": "Англ.язык",
    "Английский язык": "Англ.язык",
    "Информатика и искусственный интеллект": "Информатика",
    "Художественный труд": "Худ.труд",
    "Физическая культура": "Физкультура",
    "Основы права": "Основы права",
    "Цифровая грамотность и искусственный интеллект": "Цифр. грамотность",
    "Цифровая грамотность": "Цифр. грамотность",
    "Литературное чтение": "Лит. чтение",
    "Обучение грамоте": "Обуч. грамоте",
    "Трудовое обучение": "Труд. обучение",
    "Казахская литература": "Каз. литература",
    "Адаптивная физическая культура": "Адапт. физкультура",
    "Чтение и развитие речи": "Чтение и разв. речи",
    "Развитие речи и коммуникации": "Разв. речи и комм.",
    "Хозяйственный труд": "Хоз. труд",
    "Познавательное развитие": "Познав. развитие",
    "Индивидуально-развивающие занятия": "Инд.-развив. занятия",
    "Социально-бытовая ориентировка": "СБО",
    "Изобразительное искусство": "ИЗО",
    # Казахские предметы (если встретятся)
    "Қазақ тілі және әдебиеті": "ҚиӘ",
    "Қазақ әдебиеті": "Әдебиет",
    "Қазақ тілі мен әдебиеті": "ҚиӘ",
    "Қазақстан тарихы": "Қазақстан тарихы",
    "Дүниежүзі тарихы": "Дүниежүзі тар.",
    "Көркем еңбек": "Көркем еңбек",
    "Дене шынықтыру": "Дене шынықтыру",
    "Ағылшын тілі": "Ағылш. тілі",
    "Алгебра және анализ бастамалары": "Алгебра және анализ баст.",
    "Алғашқы әскери және технологиялық дайындық": "АӘжТД",
    "Құқық негіздері": "Құқық негіз.",
    "Жаратылыстану": "Жаратылыстану",
    "Сауат ашу": "Сауат ашу",
}


def print_subj(s: str) -> str:
    return PRINT_SHORT_SUBJECT.get(s, s)


def format_cell_subs(subs: list[str]) -> str:
    """Форматирование предметов ячейки: деление на группы англ/инфо пишется в одну строку как в бланке."""
    if not subs:
        return ""
    if len(subs) == 1:
        return subs[0]
    has_info = any("Информ" in s for s in subs)
    has_eng = any("Англ" in s or "Иностр" in s for s in subs)
    if has_info and has_eng and len(subs) == 2:
        return r"Информ\англ" if ("Информ" in subs[0]) else r"Англ\инфо"
    return " / ".join(subs)


# ---------- сбор данных: обычные классы ----------

def class_subject_cells(school, res):
    """(класс, день, урок-0based) → краткий текст предмета(ов); деление на группы компактно в ячейке."""
    cells = {}
    for mi, mt in enumerate(school.meetings):
        d, p = res["slots"][str(mi)]
        for c in mt.classes:
            subs = []
            for sec in mt.sections:
                att = [a for a in sec.attendees if a[0] == c]
                if not att:
                    continue
                text = print_subj(att[0][2])
                if text not in subs:
                    subs.append(text)
            cells[(c, d, p)] = format_cell_subs(subs)
    return cells


# ---------- сбор данных: ЛУО (включённые ученики внутри обычного класса) ----------

def luo_cells(school, res):
    """(ученик, день, урок-0based) → предмет; колонки — ЛУО-ученики, не заведённые отдельным классом."""
    data = load("luo_inclusion.json")
    columns = []
    cells = {}
    for st in data.get("students", []):
        c = st["host_class"]
        key = f"{st['student']} ({c})"
        columns.append(key)
        occ = {}  # предмет → [(день, урок0)]
        for mi, mt in enumerate(school.meetings):
            if c not in mt.classes:
                continue
            d, p = res["slots"][str(mi)]
            for sj in mt.class_subjects(c):
                occ.setdefault(sj, []).append((d, p))
        for m in st["map"]:
            subs = m.get("class_subjects") or [m["class_subject"]]
            slots = sorted({s for sj in subs for s in occ.get(sj, [])})
            attends = m.get("attends", len(slots))
            for d, p in slots[:attends]:
                text = print_subj(subs[0]) if len(subs) == 1 else "/".join(print_subj(x) for x in subs)
                cells[(key, d, p)] = text
    return columns, cells


# ---------- сбор данных: ДО (занятия на дому/индивидуально) ----------

def do_cells():
    """(ученик, день, урок-0based) → предмет; колонки — ученики ДО. Возвращает также список пропущенных записей."""
    bells = {k: v for k, v in load("time_grid.json").get("bells", {}).items()}
    students = load("special_ed_schedule.json")["students"]
    columns = []
    cells = {}
    skipped = []
    for st in students:
        key = f"{st['student']} ({st['class']})"
        columns.append(key)
        for day, items in st["schedule"].items():
            d = DAY_LABELS.index(day)
            for it in items:
                per = it["period"]
                if isinstance(per, int):
                    p = per - 1
                elif isinstance(per, str) and ":" in per:
                    a = _to_min(per)
                    hit = [n for n, span in bells.items() if int(n) <= PERIODS and
                           _to_min(span.split("-")[0]) < a + 45 and _to_min(span.split("-")[1]) > a]
                    if not hit:
                        skipped.append((key, day, it["subject"]))
                        continue
                    p = int(hit[0]) - 1
                else:
                    skipped.append((key, day, it["subject"]))
                    continue
                prev = cells.get((key, d, p))
                text = print_subj(it["subject"])
                cells[(key, d, p)] = text if not prev else f"{prev} / {text}"
    return columns, cells, skipped


# ---------- сборка листа: строки по дням, фиксированное или по факту число уроков ----------

def build_sheet(name, columns, cells, fixed_periods=None):
    days = []
    for d in range(DAYS):
        if fixed_periods is not None:
            max_p = fixed_periods
        else:
            max_p = 0
            for c in columns:
                for p in range(PERIODS):
                    if cells.get((c, d, p)):
                        max_p = max(max_p, p + 1)
            if max_p == 0:
                continue
        rows = [{"period": p + 1, "cells": [cells.get((c, d, p), "") for c in columns]}
                for p in range(max_p)]
        days.append({"label": DAY_LABELS[d], "rows": rows})
    return {"name": name, "columns": columns, "days": days}


def default_sheets(school, res, include_special: bool = True):
    classes_sorted = sorted(school.classes, key=lambda c: (grade_of(c), c))
    cells = class_subject_cells(school, res)
    primary = [c for c in classes_sorted if grade_of(c) <= 4]
    upper = [c for c in classes_sorted if grade_of(c) >= 5]
    sheets = [
        build_sheet("1-4 классы", primary, cells,
                    fixed_periods=max(school.class_max_lessons[c] for c in primary)),
        build_sheet("5-11 классы", upper, cells,
                    fixed_periods=max(school.class_max_lessons[c] for c in upper)),
    ]
    if include_special:
        from export_special_schedules import get_do_schedule, get_luo_schedule
        sheets.append(get_luo_schedule(school, res))
        sheets.append(get_do_schedule())
    return sheets


# ---------- запись в .xlsx ----------

def write_print_grid(ws, sheet, header: dict | None = None):
    header = header or HEADER_INFO
    columns = sheet["columns"]
    ncol = 2 + len(columns)

    if ncol >= 8:
        # 3-блочная компоновка по образцу фото (1-4 классы, 5-11 классы)
        left_end = 4 if ncol >= 13 else 3
        right_start = ncol - 3 if ncol >= 13 else ncol - 2
        center_start = left_end + 1
        center_end = right_start - 1

        # Строка 1: Согласовано (слева) | Утверждаю (справа)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=left_end)
        c = ws.cell(1, 1, header["agree"]["title"])
        c.font = FONT_AGREE_TITLE
        c.alignment = LEFT_ALIGN

        ws.merge_cells(start_row=1, start_column=right_start, end_row=1, end_column=ncol)
        c = ws.cell(1, right_start, header["approve"]["title"])
        c.font = FONT_AGREE_TITLE
        c.alignment = LEFT_ALIGN
        ws.row_dimensions[1].height = 17

        # Строка 2: Председатель родкомитета | Расписание уроков на I полугодие | Директор школы:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=left_end)
        c = ws.cell(2, 1, header["agree"]["role"])
        c.font = FONT_AGREE_BODY
        c.alignment = LEFT_ALIGN

        ws.merge_cells(start_row=2, start_column=center_start, end_row=2, end_column=center_end)
        c = ws.cell(2, center_start, header["title"])
        c.font = FONT_TITLE_MAIN
        c.alignment = CENTER

        ws.merge_cells(start_row=2, start_column=right_start, end_row=2, end_column=ncol)
        c = ws.cell(2, right_start, header["approve"]["role"])
        c.font = FONT_AGREE_BODY
        c.alignment = LEFT_ALIGN
        ws.row_dimensions[2].height = 20

        # Строка 3: Чоп Н.А. | КГУ «Песчанская СОШ №1» ... | Мельников П.В.
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=left_end)
        c = ws.cell(3, 1, header["agree"]["name"])
        c.font = FONT_AGREE_NAME
        c.alignment = LEFT_ALIGN

        ws.merge_cells(start_row=3, start_column=center_start, end_row=3, end_column=center_end)
        c = ws.cell(3, center_start, header["subtitle"])
        c.font = FONT_TITLE_SUB
        c.alignment = CENTER

        ws.merge_cells(start_row=3, start_column=right_start, end_row=3, end_column=ncol)
        c = ws.cell(3, right_start, header["approve"]["name"])
        c.font = FONT_AGREE_NAME
        c.alignment = LEFT_ALIGN
        ws.row_dimensions[3].height = 20

        # Строка 4: разделитель перед таблицей
        ws.row_dimensions[4].height = 8

        # Строка 5: шапка таблицы (классы)
        table_head_row = 5
        ws.cell(table_head_row, 1).border = BORDER
        ws.cell(table_head_row, 2).border = BORDER
        for i, c in enumerate(columns):
            cell = ws.cell(table_head_row, 3 + i, c)
            cell.font = BOLD
            cell.alignment = CENTER
            cell.border = BORDER
        ws.row_dimensions[table_head_row].height = 30 if any(len(c) > 12 for c in columns) else 22

        row = 6
        freeze_cell = "C6"
        print_head = "1:5"
    else:
        # Узкие листы (ЛУО, ДО) — шапка с последовательным размещением
        ws.cell(1, 1, header["agree"]["title"]).font = FONT_AGREE_TITLE
        ws.cell(1, ncol, header["approve"]["title"]).font = FONT_AGREE_TITLE
        ws.row_dimensions[1].height = 17

        ws.cell(2, 1, header["agree"]["role"]).font = FONT_AGREE_BODY
        ws.cell(2, ncol, header["approve"]["role"]).font = FONT_AGREE_BODY
        ws.row_dimensions[2].height = 18

        ws.cell(3, 1, header["agree"]["name"]).font = FONT_AGREE_NAME
        ws.cell(3, ncol, header["approve"]["name"]).font = FONT_AGREE_NAME
        ws.row_dimensions[3].height = 18

        ws.row_dimensions[4].height = 6

        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ncol)
        c = ws.cell(5, 1, header["title"])
        c.font = FONT_TITLE_MAIN
        c.alignment = CENTER
        ws.row_dimensions[5].height = 20

        sub_label = f"{header['subtitle']} ({sheet['name']})"
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=ncol)
        c = ws.cell(6, 1, sub_label)
        c.font = FONT_TITLE_SUB
        c.alignment = CENTER
        ws.row_dimensions[6].height = 18

        ws.row_dimensions[7].height = 6

        table_head_row = 8
        ws.cell(table_head_row, 1).border = BORDER
        ws.cell(table_head_row, 2).border = BORDER
        for i, c in enumerate(columns):
            cell = ws.cell(table_head_row, 3 + i, c)
            cell.font = BOLD
            cell.alignment = CENTER
            cell.border = BORDER
        ws.row_dimensions[table_head_row].height = 30 if any(len(c) > 12 for c in columns) else 22

        row = 9
        freeze_cell = "C9"
        print_head = "1:8"

    for day in sheet["days"]:
        start_row = row
        for r in day["rows"]:
            top = BORDER_DAY_TOP if r["period"] == day["rows"][0]["period"] else BORDER
            pcell = ws.cell(row, 2, r["period"])
            pcell.alignment = CENTER
            pcell.border = top
            for i, text in enumerate(r["cells"]):
                cell = ws.cell(row, 3 + i, text)
                cell.alignment = CENTER
                cell.border = top
            has_three_lines = any(x.count("\n") >= 2 for x in r["cells"])
            has_two_lines = any(x.count("\n") == 1 for x in r["cells"])
            if has_three_lines:
                ws.row_dimensions[row].height = 42
            elif has_two_lines:
                ws.row_dimensions[row].height = 30
            else:
                ws.row_dimensions[row].height = 20
            row += 1
        ws.merge_cells(start_row=start_row, start_column=1, end_row=row - 1, end_column=1)
        dcell = ws.cell(start_row, 1, day["label"])
        dcell.font = BOLD
        dcell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        for r in range(start_row, row):
            b = ws.cell(r, 1).border
            ws.cell(r, 1).border = Border(left=THIN, right=THIN,
                                          top=b.top if r == start_row else THIN,
                                          bottom=MED if r == row - 1 else THIN)

    # Подвал под таблицей
    frow = row + 1
    fcell = ws.cell(frow, 1, header["compiler"])
    fcell.font = FONT_FOOTER
    fcell.alignment = LEFT_ALIGN
    ws.merge_cells(start_row=frow, start_column=1, end_row=frow, end_column=min(6, ncol))
    ws.row_dimensions[frow].height = 22

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 5
    for i in range(len(columns)):
        col_w = 26 if ncol < 8 else (15 if len(columns) <= 12 else 12)
        ws.column_dimensions[get_column_letter(3 + i)].width = col_w

    ws.freeze_panes = freeze_cell
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = print_head
    ws.print_title_cols = "A:B"
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)


def export(school, res, path, header=None, sheets=None):
    sheets = sheets or default_sheets(school, res)
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        if not sheet["columns"] or not sheet["days"]:
            continue
        ws = wb.create_sheet(sheet["name"])
        write_print_grid(ws, sheet, header)
    wb.save(path)


def dump_json(path, header, sheets):
    path.write_text(json.dumps({"header": header, "footer": {"compiler": header["compiler"]}, "sheets": sheets},
                               ensure_ascii=False), encoding="utf-8")


def main(name: str, docx: bool = True, include_special: bool = True):
    school = build()
    res = json.loads((OUT / f"{name}.json").read_text(encoding="utf-8"))
    header = dict(HEADER_INFO)
    sheets = default_sheets(school, res, include_special=include_special)
    XLSX_DIR.mkdir(parents=True, exist_ok=True)
    path = XLSX_DIR / f"Расписание_{name}_бланк.xlsx"
    export(school, res, path, header, sheets)
    print("записан", path)
    json_path = OUT / f"{name}_print.json"
    dump_json(json_path, header, sheets)
    if not docx:
        print("данные для .docx:", json_path)
        return
    import subprocess
    docx_path = XLSX_DIR / f"Расписание_{name}_бланк.docx"
    js = Path(__file__).with_name("export_print_docx.js")
    r = subprocess.run(["node", str(js), str(json_path), str(docx_path)],
                       capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        print(r.stdout, r.stderr, sep="\n")
        raise RuntimeError("не удалось собрать .docx — см. вывод node выше")
    print("записан", docx_path)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    flags = [a for a in sys.argv[1:] if a.startswith("-")]
    name = args[0] if args else "V34"
    no_special = any(f in flags for f in ("--no-special", "--no-luo-do", "--no-luo", "-n"))
    main(name, include_special=not no_special)
