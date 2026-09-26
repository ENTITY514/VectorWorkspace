"""Экспорт расписаний ЛУО и ДО классов/учащихся в Excel (.xlsx) и Word (.docx).
Стиль школьного бланка (Песчанская СОШ №1): шапка «Согласовано» / «Утверждаю»,
тонкие чёрные рамки, альбомная ориентация.
Для ДО обязательно указывается точное время каждого урока для каждого класса!

Использование:
    py -3.13 tools/bilimclass/schedule/export_special_schedules.py [V36]
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

sys.path.insert(0, str(Path(__file__).parent))
from dataset import DAY_LABELS, DAYS, build, load
from export_print import (
    BORDER,
    BORDER_DAY_TOP,
    CENTER,
    FONT_AGREE_BODY,
    FONT_AGREE_NAME,
    FONT_AGREE_TITLE,
    FONT_FOOTER,
    FONT_TITLE_MAIN,
    FONT_TITLE_SUB,
    HEADER_INFO,
    LEFT_ALIGN,
    MED,
    THIN,
    print_subj,
)
from run_variants import OUT, XLSX_DIR


def short_fio(fio: str) -> str:
    """Сокращение ФИО: Пащенко Наталья Аркадьевна -> Пащенко Н.А."""
    parts = fio.split()
    if len(parts) >= 3:
        return f"{parts[0]} {parts[1][0]}.{parts[2][0]}."
    if len(parts) == 2:
        return f"{parts[0]} {parts[1][0]}."
    return fio


# ==============================================================================
# СБОР ДАННЫХ: ЛУО
# ==============================================================================

def get_luo_schedule(school, res) -> dict:
    """
    Собирает расписание для ЛУО:
    1) 6В класс (ЛУО-класс, 19 часов)
    2) 7А класс — Власова Анастасия (ЛУО, инклюзия в 7А, 21 час)
    3) 8Б класс — Тургумбаева Айсана (ЛУО ДО, 10 часов)
    """
    columns = [
        "6В класс\n(ЛУО)",
        "7А класс\nВласова А. (ЛУО)",
        "8Б класс\nТургумбаева А. (ЛУО ДО)",
    ]
    col_6v, col_7a, col_8b = columns

    cells: dict[tuple[str, int, int], str] = {}

    # 1. 6В класс
    for mi, mt in enumerate(school.meetings):
        if "6В" not in mt.classes:
            continue
        d, p = res["slots"][str(mi)]
        for sec in mt.sections:
            for att in sec.attendees:
                if att[0] == "6В":
                    subj = print_subj(att[2])
                    tch = short_fio(sec.teacher)
                    cells[(col_6v, d, p)] = f"{subj}\n({tch})"

    # 2. 7А класс (Власова Анастасия)
    luo_data = load("luo_inclusion.json")
    st_7a = luo_data["students"][0]
    occ_7a: dict[str, list[tuple[int, int, str]]] = {}
    for mi, mt in enumerate(school.meetings):
        if "7А" not in mt.classes:
            continue
        d, p = res["slots"][str(mi)]
        for sec in mt.sections:
            for att in sec.attendees:
                if att[0] == "7А":
                    occ_7a.setdefault(att[2], []).append((d, p, sec.teacher))

    for m in st_7a["map"]:
        subs = m.get("class_subjects") or [m["class_subject"]]
        matching = []
        for sj in subs:
            for d, p, tch in occ_7a.get(sj, []):
                matching.append((d, p, sj, tch, m["luo_subject"]))
        matching.sort()
        attends = m.get("attends", len(matching))
        for d, p, cls_sj, tch, luo_sj in matching[:attends]:
            subj_display = print_subj(luo_sj)
            tch_display = short_fio(tch)
            cells[(col_7a, d, p)] = f"{subj_display}\n({tch_display})"

    # 3. 8Б класс (Тургумбаева Айсана)
    sp_data = load("special_ed_schedule.json")
    st_8b = next(s for s in sp_data["students"] if "Тургумбаева" in s["student"])
    for day_label, items in st_8b["schedule"].items():
        d = DAY_LABELS.index(day_label)
        p3_items = [it for it in items if it["period"] == 3]
        for it in items:
            p = it["period"] - 1  # 0-based
            if p == 2 and len(p3_items) > 1:
                # Вт: СБО / История Каз.; Ср: Естествознание / География
                subjs = " / ".join(print_subj(x["subject"]) for x in p3_items)
                tchs = " / ".join(short_fio(x["teacher"]).split()[0] for x in p3_items)
                cells[(col_8b, d, p)] = f"{subjs}\n({tchs})"
                break
            else:
                subj = print_subj(it["subject"])
                tch = short_fio(it["teacher"])
                cells[(col_8b, d, p)] = f"{subj}\n({tch})"

    # Собираем строки по дням (максимум до 7 уроков)
    days = []
    for d in range(DAYS):
        max_p = 0
        for col in columns:
            for p in range(7):
                if cells.get((col, d, p)):
                    max_p = max(max_p, p + 1)
        if max_p == 0:
            continue
        rows = [
            {"period": p + 1, "cells": [cells.get((col, d, p), "") for col in columns]}
            for p in range(max_p)
        ]
        days.append({"label": DAY_LABELS[d], "rows": rows})

    return {
        "name": "ЛУО",
        "title": "Расписание уроков специального (коррекционного) образования (ЛУО)",
        "subtitle": "КГУ «Песчанская СОШ №1» 2026-2027 учебный год",
        "columns": columns,
        "days": days,
    }


# ==============================================================================
# СБОР ДАННЫХ: ДО (Надомное обучение с указанием времени для каждого класса)
# ==============================================================================

def get_do_schedule() -> dict:
    """
    Собирает расписание ДО (индивидуальное надомное обучение):
    В каждой ячейке указывается:
      1) Точное время урока (для данного класса/ученика)
      2) Предмет
      3) Учитель
    """
    columns = [
        "1Б класс\nГалле О. (ДО)",
        "7А класс\nЖелнин М. (ДО)",
        "8Б класс\nТургумбаева А. (ДО)",
        "9Б класс\nРоманов Н. (ДО)",
    ]
    col_1b, col_7a, col_8b, col_9b = columns

    cells: dict[tuple[str, int, int], str] = {}

    # --- 1Б класс: Галле Оливия ---
    # Вт: 1 ур. 08:15–09:00 Развитие речи, 2 ур. 09:10–09:50 Математика (Дукач Т.А.)
    # Ср: 2 ур. 09:10–09:50 ИЗО (0.5ч), 3 ур. 10:05–10:50 Сенсорное развитие (Квашнина Е.В.)
    # Чт: 1 ур. 08:15–09:00 Развитие речи, 2 ур. 09:10–09:50 Мир вокруг (Дукач Т.А.)
    # Пт: 2 ур. 09:10–09:50 Предметно-практ. деят., 3 ур. 10:05–10:50 Сенсорное разв., 4 ур. 11:05–11:50 Предметно-практ. деят.
    galle_items = {
        # Вт
        (1, 0): ("08:15–09:00", "Разв. речи и комм.", "Дукач Т.А."),
        (1, 1): ("09:10–09:50", "Математика", "Дукач Т.А."),
        # Ср
        (2, 1): ("09:10–09:50", "ИЗО (0.5ч)", "Квашнина Е.В."),
        (2, 2): ("10:05–10:50", "Сенсорное развитие", "Квашнина Е.В."),
        # Чт
        (3, 0): ("08:15–09:00", "Разв. речи и комм.", "Дукач Т.А."),
        (3, 1): ("09:10–09:50", "Мир вокруг", "Дукач Т.А."),
        # Пт
        (4, 1): ("09:10–09:50", "Предметно-практ. деят.", "Квашнина Е.В."),
        (4, 2): ("10:05–10:50", "Сенсорное развитие", "Квашнина Е.В."),
        (4, 3): ("11:05–11:50", "Предметно-практ. деят.", "Квашнина Е.В."),
    }
    for (d, p), (tm, subj, tch) in galle_items.items():
        cells[(col_1b, d, p)] = f"{tm}\n{subj}\n({tch})"

    # --- 7А класс: Желнин Матвей ---
    # 1 ур. 08:15–09:00, 2 ур. 09:10–09:50 каждый день
    zhelnin_items = {
        # Пн
        (0, 0): ("08:15–09:00", "Разв. речи и комм.", "Дукач Т.А."),
        (0, 1): ("09:10–09:50", "Хоз. труд", "Жукова И.М."),
        # Вт
        (1, 0): ("08:15–09:00", "ИЗО", "Жукова И.М."),
        (1, 1): ("09:10–09:50", "Человек и мир", "Дукач Т.А."),
        # Ср
        (2, 0): ("08:15–09:00", "Познав. развитие", "Жукова И.М."),
        (2, 1): ("09:10–09:50", "Счет", "Дукач Т.А."),
        # Чт
        (3, 0): ("08:15–09:00", "СБО", "Жукова И.М."),
        (3, 1): ("09:10–09:50", "Разв. речи и комм.", "Дукач Т.А."),
        # Пт
        (4, 0): ("08:15–09:00", "Счет", "Дукач Т.А."),
        (4, 1): ("09:10–09:50", "Человек и мир", "Дукач Т.А."),
    }
    for (d, p), (tm, subj, tch) in zhelnin_items.items():
        cells[(col_7a, d, p)] = f"{tm}\n{subj}\n({tch})"

    # --- 8Б класс: Тургумбаева Айсана ---
    # 2 ур. 09:10–09:50, 3 ур. 10:05–10:50 каждый день
    turg_items = {
        # Пн
        (0, 1): ("09:10–09:50", "Русский язык", "Ажикенова М.А."),
        (0, 2): ("10:05–10:50", "Математика", "Пащенко Н.А."),
        # Вт
        (1, 1): ("09:10–09:50", "Математика", "Пащенко Н.А."),
        (1, 2): ("10:05–10:50", "СБО / История Каз.", "Жукова / Мельникова"),
        # Ср
        (2, 1): ("09:10–09:50", "Лит. чтение", "Ажикенова М.А."),
        (2, 2): ("10:05–10:50", "Естествозн. / Геогр.", "Красноперова О.В."),
        # Чт
        (3, 1): ("09:10–09:50", "Казахский язык", "Асаева Ж.Х."),
        (3, 2): ("10:05–10:50", "Инд.-развив. занятия", "Мельникова Л.Р."),
        # Пт
        (4, 1): ("09:10–09:50", "Русский язык", "Ажикенова М.А."),
        (4, 2): ("10:05–10:50", "Инд.-развив. занятия", "Мельникова Л.Р."),
    }
    for (d, p), (tm, subj, tch) in turg_items.items():
        cells[(col_8b, d, p)] = f"{tm}\n{subj}\n({tch})"

    # --- 9Б класс: Романов Никита ---
    # Пн–Ср: 10:30–11:15 (3 ур.), 11:25–12:10 (4 ур.), 12:20–13:05 (5 ур.)
    # Чт: 10:05–10:50 (3 ур.), 11:05–11:50 (4 ур.)
    romanov_items = {
        # Пн
        (0, 2): ("10:30–11:15", "Разв. речи и комм.", "Дукач Т.А."),
        (0, 3): ("11:25–12:10", "Человек и мир", "Дукач Т.А."),
        # Вт
        (1, 2): ("10:30–11:15", "Счет", "Дукач Т.А."),
        (1, 3): ("11:25–12:10", "Познав. развитие", "Дукач Т.А."),
        (1, 4): ("12:20–13:05", "СБО", "Дукач Т.А."),
        # Ср
        (2, 2): ("10:30–11:15", "Разв. речи и комм.", "Дукач Т.А."),
        (2, 3): ("11:25–12:10", "Счет", "Дукач Т.А."),
        (2, 4): ("12:20–13:05", "Человек и мир", "Дукач Т.А."),
        # Чт
        (3, 2): ("10:05–10:50", "ИЗО", "Жукова И.М."),
        (3, 3): ("11:05–11:50", "Хоз. труд", "Жукова И.М."),
    }
    for (d, p), (tm, subj, tch) in romanov_items.items():
        cells[(col_9b, d, p)] = f"{tm}\n{subj}\n({tch})"

    # Собираем строки по дням (максимум до 5 уроков)
    days = []
    for d in range(DAYS):
        max_p = 0
        for col in columns:
            for p in range(5):
                if cells.get((col, d, p)):
                    max_p = max(max_p, p + 1)
        if max_p == 0:
            continue
        rows = [
            {"period": p + 1, "cells": [cells.get((col, d, p), "") for col in columns]}
            for p in range(max_p)
        ]
        days.append({"label": DAY_LABELS[d], "rows": rows})

    return {
        "name": "ДО",
        "title": "Расписание занятий индивидуального обучения на дому (ДО)",
        "subtitle": "КГУ «Песчанская СОШ №1» 2026-2027 учебный год",
        "columns": columns,
        "days": days,
    }


# ==============================================================================
# ЗАПИСЬ EXCEL (.xlsx)
# ==============================================================================

def write_special_grid(ws, sheet: dict, header: dict):
    columns = sheet["columns"]
    ncol = 2 + len(columns)

    # 1. Шапка согласования: «Согласовано» (A1:B3) | «Утверждаю» (ncol:ncol)
    ws.cell(1, 1, header["agree"]["title"]).font = FONT_AGREE_TITLE
    ws.cell(1, ncol, header["approve"]["title"]).font = FONT_AGREE_TITLE
    ws.row_dimensions[1].height = 17

    ws.cell(2, 1, header["agree"]["role"]).font = FONT_AGREE_BODY
    ws.cell(2, ncol, header["approve"]["role"]).font = FONT_AGREE_BODY
    ws.row_dimensions[2].height = 18

    ws.cell(3, 1, header["agree"]["name"]).font = FONT_AGREE_NAME
    ws.cell(3, ncol, header["approve"]["name"]).font = FONT_AGREE_NAME
    ws.row_dimensions[3].height = 18

    ws.row_dimensions[4].height = 6

    # 2. Заголовок по центру
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ncol)
    c = ws.cell(5, 1, sheet.get("title") or header["title"])
    c.font = FONT_TITLE_MAIN
    c.alignment = CENTER
    ws.row_dimensions[5].height = 22

    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=ncol)
    c = ws.cell(6, 1, sheet.get("subtitle") or header["subtitle"])
    c.font = FONT_TITLE_SUB
    c.alignment = CENTER
    ws.row_dimensions[6].height = 18

    ws.row_dimensions[7].height = 8

    # 3. Шапка таблицы (День | № | Колонки)
    table_head_row = 8
    c1 = ws.cell(table_head_row, 1, "День")
    c1.font = Font(name="Calibri", bold=True, size=10)
    c1.alignment = CENTER
    c1.border = BORDER

    c2 = ws.cell(table_head_row, 2, "№")
    c2.font = Font(name="Calibri", bold=True, size=10)
    c2.alignment = CENTER
    c2.border = BORDER

    for i, col_name in enumerate(columns):
        cell = ws.cell(table_head_row, 3 + i, col_name)
        cell.font = Font(name="Calibri", bold=True, size=10)
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[table_head_row].height = 28

    # 4. Данные сетки
    row = 9
    for day in sheet["days"]:
        start_row = row
        for r in day["rows"]:
            is_first = (r["period"] == day["rows"][0]["period"])
            top_border = BORDER_DAY_TOP if is_first else BORDER

            # № урока
            pcell = ws.cell(row, 2, r["period"])
            pcell.font = Font(name="Calibri", bold=True, size=10)
            pcell.alignment = CENTER
            pcell.border = top_border

            # Ячейки уроков
            for i, text in enumerate(r["cells"]):
                cell = ws.cell(row, 3 + i, text)
                cell.font = Font(name="Calibri", size=9)
                cell.alignment = CENTER
                cell.border = top_border

            # Высота строки: если 3 строки текста (время + предмет + учитель) -> 42pt
            has_three_lines = any(x.count("\n") >= 2 for x in r["cells"])
            has_two_lines = any(x.count("\n") == 1 for x in r["cells"])
            if has_three_lines:
                ws.row_dimensions[row].height = 42
            elif has_two_lines:
                ws.row_dimensions[row].height = 32
            else:
                ws.row_dimensions[row].height = 22

            row += 1

        # Объединение дня слева
        ws.merge_cells(start_row=start_row, start_column=1, end_row=row - 1, end_column=1)
        dcell = ws.cell(start_row, 1, day["label"])
        dcell.font = Font(name="Calibri", bold=True, size=11)
        dcell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
        for r in range(start_row, row):
            b = ws.cell(r, 1).border
            ws.cell(r, 1).border = Border(
                left=THIN, right=THIN,
                top=b.top if r == start_row else THIN,
                bottom=MED if r == row - 1 else THIN
            )

    # 5. Подвал
    frow = row + 1
    fcell = ws.cell(frow, 1, header["compiler"])
    fcell.font = FONT_FOOTER
    fcell.alignment = LEFT_ALIGN
    ws.merge_cells(start_row=frow, start_column=1, end_row=frow, end_column=ncol)
    ws.row_dimensions[frow].height = 22

    # Ширины колонок
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 5
    for i in range(len(columns)):
        col_letter = get_column_letter(3 + i)
        ws.column_dimensions[col_letter].width = 28

    # Параметры страницы: альбомная, вписать в 1 страницу по ширине
    ws.freeze_panes = "C9"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:8"
    ws.print_title_cols = "A:B"
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)


def export_workbook(sheet_list: list[dict], xlsx_path: Path, header: dict):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheet_list:
        ws = wb.create_sheet(sheet["name"])
        write_special_grid(ws, sheet, header)
    wb.save(xlsx_path)
    print("записан:", xlsx_path)


def export_docx(sheet_list: list[dict], json_path: Path, docx_path: Path, header: dict):
    payload = {
        "header": header,
        "footer": {"compiler": header["compiler"]},
        "sheets": sheet_list,
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    js = Path(__file__).with_name("export_print_docx.js")
    r = subprocess.run(
        ["node", str(js), str(json_path), str(docx_path)],
        capture_output=True,
        text=True,
        encoding="utf-8"
    )
    if r.returncode != 0:
        print(r.stdout, r.stderr, sep="\n")
        raise RuntimeError(f"не удалось собрать .docx: {docx_path}")
    print("записан:", docx_path)


# ==============================================================================
# ТОЧКА ВХОДА
# ==============================================================================

def main(version_name: str = "V36"):
    school = build()
    res_path = OUT / f"{version_name}.json"
    if not res_path.exists():
        raise FileNotFoundError(f"Файл {res_path} не найден!")
    res = json.loads(res_path.read_text(encoding="utf-8"))

    header = dict(HEADER_INFO)
    XLSX_DIR.mkdir(parents=True, exist_ok=True)

    print("Сборка расписания ЛУО...")
    luo_sheet = get_luo_schedule(school, res)

    print("Сборка расписания ДО (с точным временем)...")
    do_sheet = get_do_schedule()

    # 1. Отдельный экспорт: ЛУО
    luo_xlsx = XLSX_DIR / "Расписание_ЛУО_бланк.xlsx"
    luo_json = OUT / "luo_print.json"
    luo_docx = XLSX_DIR / "Расписание_ЛУО_бланк.docx"
    export_workbook([luo_sheet], luo_xlsx, header)
    export_docx([luo_sheet], luo_json, luo_docx, header)

    # 2. Отдельный экспорт: ДО
    do_xlsx = XLSX_DIR / "Расписание_ДО_бланк.xlsx"
    do_json = OUT / "do_print.json"
    do_docx = XLSX_DIR / "Расписание_ДО_бланк.docx"
    export_workbook([do_sheet], do_xlsx, header)
    export_docx([do_sheet], do_json, do_docx, header)

    # 3. Совмещённый экспорт: ЛУО и ДО в одной книге
    combined_xlsx = XLSX_DIR / "Расписание_ЛУО_и_ДО_бланк.xlsx"
    combined_json = OUT / "luo_do_combined_print.json"
    combined_docx = XLSX_DIR / "Расписание_ЛУО_и_ДО_бланк.docx"
    export_workbook([luo_sheet, do_sheet], combined_xlsx, header)
    export_docx([luo_sheet, do_sheet], combined_json, combined_docx, header)

    print("\nВсе специальные файлы успешно созданы:")
    print("  -", luo_xlsx)
    print("  -", luo_docx)
    print("  -", do_xlsx)
    print("  -", do_docx)
    print("  -", combined_xlsx)
    print("  -", combined_docx)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    vname = sys.argv[1] if len(sys.argv) > 1 else "V36"
    main(vname)
