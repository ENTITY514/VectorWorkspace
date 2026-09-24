"""Выгрузка варианта расписания в Excel."""
from __future__ import annotations

import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dataset import DAY_LABELS, DAYS, PERIODS, School

SHORT_SUBJECT = {
    "Информатика и искусственный интеллект": "Информатика",
    "Цифровая грамотность и искусственный интеллект": "Цифр. грамотность",
    "Цифровая грамотность": "Цифр. грамотность",
    "Казахский язык и литература": "КиЛ",
    "Русский язык и литература": "Рус. язык и лит.",
    "Начальная военная подготовка": "НВТиП",
    "Адаптивная физическая культура": "Адапт. физкультура",
    "Физическая культура": "Физкультура",
    "Изобразительное искусство": "ИЗО",
    "Художественный труд": "Худ.труд",
    "Трудовое обучение": "Труд. обучение",
    "Основы предпринимательства и бизнеса": "ОПиБ",
    "Алгебра и начала анализа": "Алгебра и начало ан.",
    "Чтение и развитие речи": "Чтение и разв. речи",
    "Иностранный язык": "Англ.язык",
    "Английский язык": "Англ.язык",
    "История Казахстана": "История Каз.",
    "Всемирная история": "Всем.история",
    "Русская литература": "Русская лит.",
    "Казахская литература": "Каз. литература",
    "Литературное чтение": "Лит. чтение",
    "Обучение грамоте": "Обуч. грамоте",
    "Социально-бытовая ориентировка": "СБО",
    "Индивидуально-развивающие занятия": "Инд.-развив. занятия",
    "Развитие речи и коммуникации": "Разв. речи и комм.",
    "Хозяйственный труд": "Хоз. труд",
    "Познавательное развитие": "Познав. развитие",
    "Основы права": "Основы права",
}

THIN = Side(style="thin", color="999999")
MED = Side(style="medium", color="333333")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL_HEAD = PatternFill("solid", fgColor="DDE4EE")
FILL_DAY = PatternFill("solid", fgColor="EEF1F5")
FILL_JOINT = PatternFill("solid", fgColor="E3F0FF")
FILL_SPLIT = PatternFill("solid", fgColor="FFF6DA")
FILL_FIXED = PatternFill("solid", fgColor="F0E6FA")
FILL_WARN = PatternFill("solid", fgColor="FDE2E1")
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")
WRAP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")


def short_name(fio: str) -> str:
    parts = fio.split()
    if len(parts) >= 3:
        return f"{parts[0]} {parts[1][0]}.{parts[2][0]}."
    if len(parts) == 2:
        return f"{parts[0]} {parts[1][0]}."
    return fio


def room_label(school: School, rid: str | None) -> str:
    if not rid:
        return ""
    r = school.rooms.get(rid)
    if not r:
        return rid
    name = r["name"]
    name = re.sub(r"^Кабинет\s+", "каб. ", name)
    name = re.sub(r"^Начальный класс\s+", "нач. класс ", name)
    name = re.sub(r"^Спортивный зал", "спортзал", name)
    return name[0].lower() + name[1:] if name.startswith("Малый") else name


def subj(s: str) -> str:
    return SHORT_SUBJECT.get(s, s)


def class_cells(school: School, res: dict):
    """(класс, день, урок) → (текст, заливка)."""
    cells = {}
    for mi, mt in enumerate(school.meetings):
        d, p = res["slots"][str(mi)]
        for c in mt.classes:
            lines = []
            multi = len(mt.sections) > 1
            for si, sec in enumerate(mt.sections):
                att = [a for a in sec.attendees if a[0] == c]
                if not att:
                    continue
                room = room_label(school, res["rooms"][f"{mi}|{si}"])
                sub = att[0][1]
                text = subj(att[0][2])
                if multi and sub:
                    lines.append(f"{sub}: {text}\n{short_name(sec.teacher)}, {room}")
                else:
                    lines.append(f"{text}\n{short_name(sec.teacher)}, {room}")
            others = sorted(mt.classes - {c})
            if others:
                lines.append("вместе с " + ", ".join(others))
            fill = FILL_JOINT if others else (FILL_SPLIT if multi else None)
            cells[(c, d, p)] = ("\n".join(lines), fill)
    return cells


def aisana_cells(school: School):
    a = school.aisana
    out = defaultdict(list)
    for di, day in enumerate(DAY_LABELS):
        for item in a["schedule"].get(day, []):
            out[(di, item["period"] - 1)].append(f"{subj(item['subject'])}\n{short_name(item['teacher'])}")
    return {k: ("\n— или —\n".join(v) + ("\n(через неделю)" if len(v) > 1 else ""), FILL_FIXED) for k, v in out.items()}


def write_class_sheet(wb, school: School, res: dict, title: str):
    ws = wb.active
    ws.title = "Классы"
    cells = class_cells(school, res)
    aisana = aisana_cells(school)
    columns = []
    for c in school.classes:
        label = c + ("\nЛУО" if school.class_type.get(c) == "luo" else "")
        if c == "7А":
            label += "\n(+ инклюзия ЛУО)"
        columns.append((label, lambda d, p, c=c: cells.get((c, d, p))))
        if c == "8Б":
            columns.append((f"8Б ДО\n{short_name(school.aisana['student'])}\n(на дому, фикс.)",
                            lambda d, p: aisana.get((d, p))))

    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(columns))
    hr = 2
    ws.cell(hr, 1, "День").fill = FILL_HEAD
    ws.cell(hr, 2, "Урок").fill = FILL_HEAD
    for j, (label, _) in enumerate(columns):
        cell = ws.cell(hr, 3 + j, label)
        cell.fill = FILL_HEAD
        cell.font = Font(bold=True, size=11)
    for col in range(1, 3 + len(columns)):
        ws.cell(hr, col).alignment = WRAP
        ws.cell(hr, col).border = BORDER
    ws.row_dimensions[hr].height = 48

    row = hr + 1
    for d in range(DAYS):
        start = row
        for p in range(PERIODS):
            ws.cell(row, 2, p + 1).alignment = WRAP
            ws.cell(row, 2).fill = FILL_DAY
            ws.cell(row, 2).border = BORDER
            for j, (_, getter) in enumerate(columns):
                v = getter(d, p)
                cell = ws.cell(row, 3 + j, v[0] if v else None)
                cell.alignment = WRAP
                cell.border = BORDER
                cell.font = Font(size=9)
                if v and v[1]:
                    cell.fill = v[1]
            ws.row_dimensions[row].height = 58
            row += 1
        ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)
        dc = ws.cell(start, 1, DAY_LABELS[d])
        dc.alignment = Alignment(text_rotation=90, vertical="center", horizontal="center")
        dc.font = Font(bold=True, size=14)
        dc.fill = FILL_DAY
        for r_ in range(start, row):
            ws.cell(r_, 1).border = BORDER
        for col in range(1, 3 + len(columns)):
            c_ = ws.cell(row - 1, col)
            c_.border = Border(left=c_.border.left, right=c_.border.right, top=c_.border.top, bottom=MED)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 5
    for j in range(len(columns)):
        ws.column_dimensions[get_column_letter(3 + j)].width = 21
    ws.freeze_panes = ws.cell(hr + 1, 3)
    legend = row + 1
    ws.cell(legend, 3, "Обозначения:").font = Font(bold=True)
    for k, (fill, text) in enumerate([(FILL_JOINT, "класс-комплект (несколько классов вместе)"),
                                      (FILL_SPLIT, "деление на подгруппы"),
                                      (FILL_FIXED, "фиксированные занятия ДО (на дому)")]):
        ws.cell(legend + 1 + k, 3).fill = fill
        ws.cell(legend + 1 + k, 4, text)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_teacher_sheet(wb, school: School, res: dict):
    ws = wb.create_sheet("Учителя")
    grid = defaultdict(dict)
    for mi, mt in enumerate(school.meetings):
        d, p = res["slots"][str(mi)]
        for si, sec in enumerate(mt.sections):
            cls = ",".join(sorted({a[0] + (f"/{a[1]}" if a[1] else "") for a in sec.attendees}))
            grid[sec.teacher][(d, p)] = f"{cls}\n{room_label(school, res['rooms'][f'{mi}|{si}'])}"
    ws.cell(1, 1, "Учитель").fill = FILL_HEAD
    ws.cell(2, 1).fill = FILL_HEAD
    for d in range(DAYS):
        c0 = 2 + d * PERIODS
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + PERIODS - 1)
        ws.cell(1, c0, DAY_LABELS[d]).alignment = WRAP
        ws.cell(1, c0).fill = FILL_HEAD
        ws.cell(1, c0).font = Font(bold=True)
        for p in range(PERIODS):
            cell = ws.cell(2, c0 + p, p + 1)
            cell.alignment = WRAP
            cell.fill = FILL_HEAD
            cell.border = BORDER
    ws.cell(1, 2 + DAYS * PERIODS, "Уроков").fill = FILL_HEAD
    ws.cell(1, 3 + DAYS * PERIODS, "Окон").fill = FILL_HEAD
    for i, t in enumerate(sorted(grid)):
        r = 3 + i
        ws.cell(r, 1, short_name(t)).border = BORDER
        windows = 0
        for d in range(DAYS):
            ps = sorted(p for (dd, p) in grid[t] if dd == d)
            gaps = {g for g in range(ps[0], ps[-1] + 1) if g not in ps and school.availability[t][d][g]} if ps else set()
            windows += len(gaps)
            for p in range(PERIODS):
                cell = ws.cell(r, 2 + d * PERIODS + p, grid[t].get((d, p)))
                cell.alignment = WRAP
                cell.border = BORDER
                cell.font = Font(size=8)
                if p in gaps:
                    cell.fill = FILL_WARN
                    cell.value = "окно"
                if not school.availability[t][d][p] and (d, p) not in grid[t]:
                    cell.fill = PatternFill("solid", fgColor="D9D9D9")
        ws.cell(r, 2 + DAYS * PERIODS, len(grid[t])).border = BORDER
        ws.cell(r, 3 + DAYS * PERIODS, windows).border = BORDER
        ws.row_dimensions[r].height = 30
    ws.column_dimensions["A"].width = 20
    for j in range(DAYS * PERIODS):
        ws.column_dimensions[get_column_letter(2 + j)].width = 7.5
    ws.freeze_panes = "B3"
    n = 4 + len(grid)
    ws.cell(n, 1, "Серым — учитель недоступен (совместитель / занятия ДО); красным — окно.").font = Font(italic=True)


def write_room_sheet(wb, school: School, res: dict):
    ws = wb.create_sheet("Кабинеты")
    grid = defaultdict(dict)
    for mi, mt in enumerate(school.meetings):
        d, p = res["slots"][str(mi)]
        for si, sec in enumerate(mt.sections):
            rid = res["rooms"][f"{mi}|{si}"]
            cls = ",".join(sorted({a[0] for a in sec.attendees}))
            grid[rid][(d, p)] = f"{cls}\n{short_name(sec.teacher).split()[0]}"
    ws.cell(1, 1, "Кабинет").fill = FILL_HEAD
    for d in range(DAYS):
        c0 = 2 + d * PERIODS
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + PERIODS - 1)
        ws.cell(1, c0, DAY_LABELS[d]).fill = FILL_HEAD
        for p in range(PERIODS):
            ws.cell(2, c0 + p, p + 1).fill = FILL_HEAD
    for i, rid in enumerate(sorted(grid, key=lambda r: school.rooms[r]["name"])):
        r = 3 + i
        ws.cell(r, 1, school.rooms[rid]["name"]).alignment = WRAP_LEFT
        for (d, p), v in grid[rid].items():
            cell = ws.cell(r, 2 + d * PERIODS + p, v)
            cell.alignment = WRAP
            cell.font = Font(size=8)
            cell.border = BORDER
        ws.row_dimensions[r].height = 30
    ws.column_dimensions["A"].width = 30
    for j in range(DAYS * PERIODS):
        ws.column_dimensions[get_column_letter(2 + j)].width = 7.5
    ws.freeze_panes = "B3"


def write_report_sheet(wb, lines: list[tuple[str, str]]):
    ws = wb.create_sheet("Проверка и качество")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 110
    for i, (k, v) in enumerate(lines):
        a = ws.cell(i + 1, 1, k)
        b = ws.cell(i + 1, 2, v)
        a.alignment = WRAP_LEFT
        b.alignment = WRAP_LEFT
        if k and not v:
            a.font = Font(bold=True, size=12)
        lines_n = max(1, str(v).count("\n") + 1, len(str(v)) // 100 + 1)
        ws.row_dimensions[i + 1].height = min(400, 15 * lines_n)


def _to_min(hhmm: str) -> int:
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def _add_minutes(hhmm: str, delta: int) -> str:
    h, m = (int(x) for x in hhmm.split(":"))
    total = h * 60 + m + delta
    return f"{total // 60:02d}:{total % 60:02d}"


def write_do_sheet(wb):
    """Отдельный лист с расписанием учеников на дому (ДО) — день, урок, время, предмет, учитель."""
    from dataset import load
    ws = wb.create_sheet("Расписание ДО")
    bells = load("time_grid.json").get("bells", {})
    students = load("special_ed_schedule.json")["students"]
    ws.cell(1, 1, "Расписание учеников на дому (ДО)").font = Font(bold=True, size=14)
    note = ("Время: если урок указан номером — по звонкам школы; если временем — как в документе ученика "
            "(занятие 45 минут). Учителя в эти часы заняты и в школе на уроки не ставятся.")
    ws.cell(2, 1, note).alignment = WRAP_LEFT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
    ws.row_dimensions[2].height = 30
    row = 4
    for st in students:
        head = f"{st['student']} — {st['class']}, {st.get('program', 'ДО')}"
        c = ws.cell(row, 1, head)
        c.font = Font(bold=True, size=12)
        c.fill = FILL_HEAD
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1
        extra = "  ".join(x for x in (st.get("time_note"), st.get("change_note")) if x)
        if extra:
            ws.cell(row, 1, extra).font = Font(italic=True, size=9)
            ws.cell(row, 1).alignment = WRAP_LEFT
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
            ws.row_dimensions[row].height = 28
            row += 1
        for j, h in enumerate(["День", "Урок", "Время по звонкам", "Время по документу", "Предмет",
                                "Учитель", "Примечание"]):
            cell = ws.cell(row, 1 + j, h)
            cell.font = Font(bold=True, size=10)
            cell.fill = FILL_DAY
            cell.border = BORDER
        row += 1
        for day in DAY_LABELS:
            items = st["schedule"].get(day, [])
            if not items:
                vals = [day, "—", "", "", "занятий нет", "", ""]
                for j, v in enumerate(vals):
                    cell = ws.cell(row, 1 + j, v)
                    cell.border = BORDER
                    cell.font = Font(size=10, color="888888")
                row += 1
                continue
            for it in items:
                per = it["period"]
                if isinstance(per, int):
                    lesson, time = f"{per}-й", bells.get(str(per), "").replace("-", "–")
                elif ":" in str(per):
                    a = _to_min(per)
                    hit = [n for n, span in bells.items() if int(n) <= PERIODS and
                           _to_min(span.split("-")[0]) < a + 45 and _to_min(span.split("-")[1]) > a]
                    lesson = "–".join(sorted(set([hit[0], hit[-1]]), key=int)) if hit else "—"
                    time = "; ".join(f"{n}-й {bells[n].replace('-', '–')}" for n in hit)
                else:
                    lesson, time = "—", "уточняется"
                remark = it.get("note", "")
                if it.get("hours"):
                    remark = (f"{it['hours']} ч; " + remark).strip("; ")
                vals = [day, lesson, time, it.get("doc_time", ""), it["subject"], it["teacher"], remark]
                for j, v in enumerate(vals):
                    cell = ws.cell(row, 1 + j, v)
                    cell.border = BORDER
                    cell.alignment = WRAP_LEFT
                    cell.font = Font(size=10)
                row += 1
        row += 1
    for col, w in zip("ABCDEFG", [8, 8, 26, 16, 34, 30, 34]):
        ws.column_dimensions[col].width = w


def export(school: School, res: dict, path, title: str, report_lines: list[tuple[str, str]]):
    wb = Workbook()
    write_class_sheet(wb, school, res, title)
    write_teacher_sheet(wb, school, res)
    write_room_sheet(wb, school, res)
    write_report_sheet(wb, report_lines)
    write_do_sheet(wb)
    wb.save(path)
