import json
import os
import sys
import re
from pathlib import Path
from collections import defaultdict

# Ensure UTF-8
if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

root = Path(__file__).parent.parent
catalog_path = root / "data" / "synthetic" / "catalog.json"

try:
    import openpyxl
except ImportError:
    print("openpyxl is required. Please install openpyxl.")
    sys.exit(1)


def parse_schedule_excel(excel_path: Path):
    print(f"Loading Excel file from: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 1. Parse Rooms Sheet
    rooms = []
    if "🏫 Кабинеты" in wb.sheetnames:
        ws_r = wb["🏫 Кабинеты"]
        for row in ws_r.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            idx, rid, name, rtype, teacher = (list(row) + [None]*5)[:5]
            rooms.append({
                "id": str(rid).strip(),
                "name": str(name).strip() if name else str(rid),
                "room_type": str(rtype).strip() if rtype else "General",
                "base_teacher_name": str(teacher).strip() if teacher and teacher != "—" else None
            })
    print(f"Parsed {len(rooms)} rooms.")

    # 2. Parse Teachers Sheet
    teachers = []
    if "👨‍🏫 Учителя" in wb.sheetnames:
        ws_t = wb["👨‍🏫 Учителя"]
        for row in ws_t.iter_rows(min_row=2, values_only=True):
            if not row or not row[1] or "[Впишите" in str(row[1]):
                continue
            idx, full_name, subjs, broom, max_d, is_comb, arrival_note = (list(row) + [None]*7)[:7]
            teachers.append({
                "id": re.sub(r'[^a-zа-яәіңғқөұүһ0-9]+', '_', str(full_name).lower().replace('ё', 'е')).strip('_'),
                "full_name": str(full_name).strip(),
                "base_room_name": str(broom).strip() if broom else None,
                "max_daily_lessons": int(max_d) if max_d else 7,
                "is_combined": str(is_comb).strip().lower() == "да",
                "arrival_note": str(arrival_note).strip() if arrival_note else ""
            })
    print(f"Parsed {len(teachers)} teachers.")

    # 3. Parse Classes Sheet (NO HEADCOUNT)
    classes = []
    if "👥 Классы" in wb.sheetnames:
        ws_c = wb["👥 Классы"]
        for row in ws_c.iter_rows(min_row=2, values_only=True):
            if not row or not row[1]:
                continue
            idx, cid, grade, letter, ctype, shift, note = (list(row) + [None]*7)[:7]
            classes.append({
                "id": str(cid).strip(),
                "grade": int(grade) if grade else 1,
                "letter": str(letter).strip() if letter else "",
                "type": str(ctype).strip() if ctype else "normal",
                "shift": "First" if "1" in str(shift) else "Second"
            })
    print(f"Parsed {len(classes)} classes.")

    # 4. Parse Curriculum Sheet
    curriculum = []
    if "📚 Учебный план (Нагрузка)" in wb.sheetnames:
        ws_cur = wb["📚 Учебный план (Нагрузка)"]
        for row in ws_cur.iter_rows(min_row=2, values_only=True):
            if not row or not row[1] or not row[2] or "[Новый предмет]" in str(row[2]):
                continue
            # Header or section row
            if str(row[0]).startswith("Класс ") and not row[3]:
                continue
            
            idx, cid, subj, hrs, t1, is_split, t2, joint, room_type, note = (list(row) + [None]*10)[:10]
            if not cid or not subj or not hrs:
                continue
            
            try:
                hours = int(hrs)
            except ValueError:
                continue
            
            if hours <= 0:
                continue

            t1_str = str(t1).strip() if t1 else ""
            t2_str = str(t2).strip() if t2 else ""
            
            t1_id = re.sub(r'[^a-zа-яәіңғқөұүһ0-9]+', '_', t1_str.lower().replace('ё', 'е')).strip('_') if t1_str else None
            t2_id = re.sub(r'[^a-zа-яәіңғқөұүһ0-9]+', '_', t2_str.lower().replace('ё', 'е')).strip('_') if t2_str else None
            
            s_slug = re.sub(r'[^a-zа-яәіңғқөұүһ0-9]+', '_', str(subj).lower().replace('ё', 'е')).strip('_')
            j_slug = re.sub(r'[^a-zа-яәіңғқөұүһ0-9]+', '_', str(joint).lower().replace('ё', 'е')).strip('_') if joint else None

            curriculum.append({
                "class_id": str(cid).strip(),
                "subject_name": str(subj).strip(),
                "subject_id": s_slug,
                "hours_per_week": hours,
                "teacher_name": t1_str,
                "teacher_id": t1_id,
                "split_teacher2_name": t2_str,
                "split_teacher2_id": t2_id,
                "joint_lesson_id": j_slug,
                "room_type": str(room_type).strip() if room_type else "Стандартный",
                "note": str(note).strip() if note else ""
            })
    print(f"Parsed {len(curriculum)} curriculum entries.")

    return {
        "classes": classes,
        "teachers": teachers,
        "rooms": rooms,
        "curriculum": curriculum
    }

if __name__ == "__main__":
    excel_file = root / "Materials" / "Учебный_план_и_нагрузка_2026-2027.xlsx"
    if not excel_file.exists():
        print(f"File {excel_file} not found. Run generate_schedule_template_excel.py first.")
        sys.exit(1)
    data = parse_schedule_excel(excel_file)
    print("\nSummary:")
    print(f"  Classes: {len(data['classes'])}")
    print(f"  Teachers: {len(data['teachers'])}")
    print(f"  Rooms: {len(data['rooms'])}")
    print(f"  Curriculum Entries: {len(data['curriculum'])}")
