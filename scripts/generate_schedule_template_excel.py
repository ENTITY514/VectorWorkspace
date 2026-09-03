import json
import os
import sys
from pathlib import Path

# Ensure UTF-8
if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

root = Path(__file__).parent.parent
catalog_path = root / "data" / "synthetic" / "catalog.json"
curr_path = root / "data" / "synthetic" / "curriculum_q4_2026_variant.json"
output_path = root / "Materials" / "Учебный_план_и_нагрузка_2026-2027.xlsx"

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    print("openpyxl is required. Please install openpyxl.")
    sys.exit(1)

with open(catalog_path, "r", encoding="utf-8") as f:
    catalog = json.load(f)

with open(curr_path, "r", encoding="utf-8") as f:
    curriculum = json.load(f)

# ─── 1. TEACHERS DATA ───
teachers_raw = [t for t in catalog["teachers"] if "давранова" not in t["id"].lower() and "давранова" not in t["full_name"].lower()]
teachers_list = sorted(teachers_raw, key=lambda x: x["full_name"])

# ─── 2. ROOMS DATA ───
rooms_data = [
    {"index": 1, "id": "r_1_general", "name": "Кабинет 1 (Начальные классы)", "type": "Стандартный", "teacher": "Начальные классы 1"},
    {"index": 2, "id": "r_2_gym", "name": "Кабинет 2 (Спортивный зал)", "type": "Спортзал", "teacher": "Ольков Г.В."},
    {"index": 3, "id": "r_3_general", "name": "Кабинет 3 (История)", "type": "Стандартный", "teacher": "Досаева М.Х."},
    {"index": 4, "id": "r_4_lang", "name": "Кабинет 4 (Английский язык)", "type": "Языковой", "teacher": "Акауова И.Н."},
    {"index": 5, "id": "r_5_general", "name": "Кабинет 5 (Начальные классы)", "type": "Стандартный", "teacher": "Начальные классы 2"},
    {"index": 6, "id": "r_6_general", "name": "Кабинет 6 (Казахская литература)", "type": "Стандартный", "teacher": "Ізбасты М.Ж."},
    {"index": 7, "id": "r_7_physics", "name": "Кабинет 7 (Физика)", "type": "Физика / Лаборатория", "teacher": "Русина Н.В."},
    {"index": 8, "id": "r_8_chemistry", "name": "Кабинет 8 (Химия)", "type": "Химия / Лаборатория", "teacher": "Бабич И.Д."},
    {"index": 9, "id": "r_9_nvp", "name": "Кабинет 9 (НВП)", "type": "Стандартный / НВП", "teacher": "Жолдаспаев А.К."},
    {"index": 10, "id": "r_10_informatics", "name": "Кабинет 10 (Информатика)", "type": "Информатика", "teacher": "Чифин Д.Н."},
    {"index": 11, "id": "r_11_robotics", "name": "Кабинет 11 (Робототехника)", "type": "Робототехника", "teacher": "—"},
    {"index": 12, "id": "r_12_lang", "name": "Кабинет 12 (Казахский язык)", "type": "Языковой", "teacher": "Жекеева Н.А."},
    {"index": 13, "id": "r_13_lang", "name": "Кабинет 13 (Иностранный язык)", "type": "Языковой", "teacher": "Алекпарова А.Ж."},
    {"index": 14, "id": "r_14_general", "name": "Кабинет 14 (Начальные классы)", "type": "Стандартный", "teacher": "Начальные классы 3"},
    {"index": 15, "id": "r_15_general", "name": "Кабинет 15 (Начальные классы)", "type": "Стандартный", "teacher": "Начальные классы 4"},
    {"index": 16, "id": "r_16_math", "name": "Кабинет 16 (Математика)", "type": "Стандартный", "teacher": "Пащенко Н.А."},
    {"index": 17, "id": "r_17_lang", "name": "Кабинет 17 (Русский язык)", "type": "Языковой", "teacher": "Дмитриев Е.В."},
    {"index": 18, "id": "r_18_biology", "name": "Кабинет 18 (Биология)", "type": "Биология / Лаборатория", "teacher": "Красноперова О.В."},
    {"index": 19, "id": "r_19_general", "name": "Кабинет 19 (Общий)", "type": "Стандартный", "teacher": "—"},
    {"index": 20, "id": "r_20_general", "name": "Кабинет 20 (Общий)", "type": "Стандартный", "teacher": "—"},
    {"index": 21, "id": "r_21_workshop", "name": "Кабинет 21 (Мастерская / Труд)", "type": "Мастерская", "teacher": "Сидоров В.Б."},
    {"index": 22, "id": "r_22_workshop", "name": "Кабинет 22 (Мастерская / Труд)", "type": "Мастерская", "teacher": "—"},
    {"index": 23, "id": "r_23_general", "name": "Кабинет 23 (Общий)", "type": "Стандартный", "teacher": "—"},
    {"index": 24, "id": "r_24_general", "name": "Кабинет 24 (Общий)", "type": "Стандартный", "teacher": "—"},
]

# ─── 3. CLASSES FOR NEW YEAR (+1 YEAR SHIFT) ───
class_shift_map = {
    "1а": "2а",
    "1б": "2б",
    "2": "3",
    "3а": "4а",
    "3б": "4б",
    "3_luo": "4 ЛУО",
    "4а": "5а",
    "4б": "5б",
    "4_do": "5 ДО",
    "5а": "6а",
    "5б": "6б",
    "5_luo": "6 ЛУО",
    "6": "7",
    "6_luo": "7 ЛУО",
    "6_do": "7 ДО",
    "7а": "8а",
    "7ә": "8ә",
    "7б": "8б",
    "7б_luo": "8б ЛУО",
    "8": "9",
    "8_do": "9 ДО",
    "9": "10",
    "9_do": "10 ДО",
    "10а": "11а",
    "10б": "11б",
}

new_classes_info = [
    {"id": "1а", "grade": 1, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Новый набор"},
    {"id": "1б", "grade": 1, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Новый набор"},
    {"id": "2а", "grade": 2, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Переход из 1а"},
    {"id": "2б", "grade": 2, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Переход из 1б"},
    {"id": "3", "grade": 3, "letter": "", "type": "Обычный", "shift": "1 смена", "note": "Переход из 2"},
    {"id": "4а", "grade": 4, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Переход из 3а"},
    {"id": "4б", "grade": 4, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Переход из 3б"},
    {"id": "4 ЛУО", "grade": 4, "letter": "", "type": "ЛУО (спецкласс)", "shift": "1 смена", "note": "Переход из 3 ЛУО"},
    {"id": "5а", "grade": 5, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Переход из 4а"},
    {"id": "5б", "grade": 5, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Переход из 4б"},
    {"id": "5 ДО", "grade": 5, "letter": "", "type": "Домашнее обучение", "shift": "1 смена", "note": "Переход из 4 ДО"},
    {"id": "6а", "grade": 6, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Переход из 5а"},
    {"id": "6б", "grade": 6, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Переход из 5б"},
    {"id": "6 ЛУО", "grade": 6, "letter": "", "type": "ЛУО (спецкласс)", "shift": "1 смена", "note": "Переход из 5 ЛУО"},
    {"id": "7", "grade": 7, "letter": "", "type": "Обычный", "shift": "1 смена", "note": "Переход из 6"},
    {"id": "7 ЛУО", "grade": 7, "letter": "", "type": "ЛУО (спецкласс)", "shift": "1 смена", "note": "Переход из 6 ЛУО"},
    {"id": "7 ДО", "grade": 7, "letter": "", "type": "Домашнее обучение", "shift": "1 смена", "note": "Переход из 6 ДО"},
    {"id": "8а", "grade": 8, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Переход из 7а"},
    {"id": "8ә", "grade": 8, "letter": "ә", "type": "Обычный (каз)", "shift": "1 смена", "note": "Переход из 7ә"},
    {"id": "8б", "grade": 8, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Переход из 7б"},
    {"id": "8б ЛУО", "grade": 8, "letter": "б", "type": "ЛУО (спецкласс)", "shift": "1 смена", "note": "Переход из 7б ЛУО"},
    {"id": "9", "grade": 9, "letter": "", "type": "Обычный", "shift": "1 смена", "note": "Переход из 8"},
    {"id": "9 ДО", "grade": 9, "letter": "", "type": "Домашнее обучение", "shift": "1 смена", "note": "Переход из 8 ДО"},
    {"id": "10", "grade": 10, "letter": "", "type": "Обычный", "shift": "1 смена", "note": "Переход из 9"},
    {"id": "10 ДО", "grade": 10, "letter": "", "type": "Домашнее обучение", "shift": "1 смена", "note": "Переход из 9 ДО"},
    {"id": "11а", "grade": 11, "letter": "а", "type": "Обычный", "shift": "1 смена", "note": "Переход из 10а"},
    {"id": "11б", "grade": 11, "letter": "б", "type": "Обычный", "shift": "1 смена", "note": "Переход из 10б"},
]

# ─── 4. SUBJECT HUMAN NAMES & SPECIAL ROOM TYPES ───
subject_display_names = {
    "pe": "Физическая культура",
    "bukvar": "Обучение грамоте",
    "music": "Музыка",
    "math": "Математика",
    "algebra": "Алгебра",
    "geometry": "Геометрия",
    "estestvoznanie": "Естествознание",
    "trud": "Трудовое обучение / Труд",
    "izo": "Изобразительное искусство",
    "kazakh_language": "Казахский язык",
    "kazakh_literature": "Казахская литература",
    "russian": "Русский язык",
    "literature": "Русская литература",
    "english": "Иностранный (Английский) язык",
    "informatics": "Информатика",
    "physics": "Физика",
    "chemistry": "Химия",
    "biology": "Биология",
    "geography": "География",
    "history_kz": "История Казахстана",
    "history_world": "Всемирная история",
    "law": "Основы права",
    "nvp": "НВП / НВиТП",
    "poznanie": "Познание мира",
    "sbo": "СБО (Социально-бытовая ориентировка)",
    "ch_rr": "Чтение и развитие речи",
    "pto": "ПТО (Профессионально-трудовое обучение)",
    "schet": "Счет",
    "afk": "АФК (Адаптивная физкультура)",
    "к_ркем_е_бек": "Көркем еңбек",
    "орыс_т_л": "Орыс тілі",
    "орыс_т_л_ж_е_деб": "Орыс тілі және әдебиеті",
    "д_ниеж_з_тарихы": "Дүниежүзі тарихы",
    "афк": "АФК (Адаптивная физкультура)",
    "цг": "ЦГ",
    "цс": "ЦС",
    "нвитп": "НВиТП",
    "ррик": "РРиК",
}

def get_subj_name(sid):
    if not sid:
        return ""
    if sid in subject_display_names:
        return subject_display_names[sid]
    for s in catalog["subjects"]:
        if s["id"] == sid:
            return s["name"]
    return sid

def get_required_room_type(subj_name, sid=""):
    s_low = (subj_name + " " + sid).lower()
    if "физ" in s_low or "дене" in s_low or "pe" in s_low:
        return "Спортзал"
    if "информ" in s_low:
        return "Информатика"
    if "робот" in s_low:
        return "Робототехника"
    if "хим" in s_low:
        return "Химия / Лаборатория"
    if "физик" in s_low:
        return "Физика / Лаборатория"
    if "биолог" in s_low:
        return "Биология / Лаборатория"
    if "труд" in s_low or "еңбек" in s_low or "мастер" in s_low or "пто" in s_low:
        return "Мастерская"
    if "казах" in s_low or "қазақ" in s_low or "англ" in s_low or "ағыл" in s_low or "немец" in s_low or "язык" in s_low:
        return "Языковой"
    return "Стандартный"

def is_split_subj(subj_name, sid=""):
    s_low = (subj_name + " " + sid).lower()
    if "казах" in s_low or "қазақ" in s_low or "англ" in s_low or "ағыл" in s_low or "немец" in s_low:
        return "Да"
    if "информ" in s_low:
        return "Да"
    if "труд" in s_low or "еңбек" in s_low or "көркем" in s_low:
        return "Да"
    return "Нет"

teacher_name_by_id = {t["id"]: t["full_name"] for t in catalog["teachers"]}

print("Generating clean Excel Workbook with multi-class joint lessons support...")
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Setup styles
font_title = Font(name="Calibri", size=14, bold=True, color="1F4E79")
font_section = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=10)
font_bold = Font(name="Calibri", size=10, bold=True)
font_muted = Font(name="Calibri", size=9, italic=True, color="595959")

fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
fill_blue = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
fill_light_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
fill_zebra = PatternFill(start_color="F9FBFD", end_color="F9FBFD", fill_type="solid")
fill_accent = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
fill_empty_row = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")

thin_border_side = Side(style="thin", color="D9D9D9")
thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# ═══════════════════════════════════════════════════════════════════════════
# SHEET 1: 📌 ИНСТРУКЦИЯ
# ═══════════════════════════════════════════════════════════════════════════
ws_info = wb.create_sheet(title="📌 Инструкция")
ws_info.views.sheetView[0].showGridLines = True

info_texts = [
    ("Инструкция по заполнению Учебного Плана и Нагрузки на 2026–2027 учебный год", font_title),
    ("", font_data),
    ("Данная таблица предназначена для распределения учебных часов, закрепления учителей и формирования расписания в VectorWorkspace.", font_data),
    ("", font_data),
    ("📋 1. ОСНОВНОЙ ЛИСТ «Учебный план (Нагрузка)»:", font_bold),
    ("  • Все предметы и классы уже предварительно заполнены с учетом перехода (+1 год к параллели).", font_data),
    ("  • В начальных классах (1–4 классы) учителя оставлены пустыми — впишите ФИО учителей начальной школы.", font_data),
    ("  • В средних и старших классах (5–11 классы) учителя сохранены с прошлого года по принципу преемственности.", font_data),
    ("  • Для каждого класса внизу оставлены 3 пустые розовые строки для добавления новых предметов.", font_data),
    ("", font_data),
    ("👥 2. ДЕЛЕНИЕ НА ПОДГРУППЫ И ВТОРОЙ УЧИТЕЛЬ:", font_bold),
    ("  • Колонка «Деление на подгруппы»: выберите «Да» для предметов, делящихся на 2 группы (Языки, Информатика, Труд).", font_data),
    ("  • Колонка «Второй учитель»: обязательна к заполнению ТОЛЬКО если выбрано «Да». Это педагог, который ведет вторую подгруппу в отдельном кабинете.", font_data),
    ("  • Урок для обеих подгрупп встанет в один и тот же день и период параллельно.", font_data),
    ("", font_data),
    ("🔗 3. КЛАСС-КОМПЛЕКТЫ (ОБЪЕДИНЕНИЕ 2, 3 И БОЛЕЕ КЛАССОВ):", font_bold),
    ("  • Класс-комплект позволяет объединять ЛЮБОЕ количество классов на один урок (например, 2 класса: 8б + 8б ЛУО; или 3 класса: 1а + 2а + 4а на Музыку).", font_data),
    ("  • Чтобы объединить уроки нескольких классов: выберите ОДНО И ТО ЖЕ название класс-комплекта в строках этих классов.", font_data),
    ("  • Солвер автоматически поставит эти уроки в один день, один период и в один кабинет, а учителю засчитает 1 физический урок.", font_data),
    ("", font_data),
    ("👨‍🏫 4. СПИСОК УЧИТЕЛЕЙ И КАБИНЕТЫ:", font_bold),
    ("  • Лист «Учителя»: проверьте список, укажите совместителей и дни их приезда. Лимит уроков по умолчанию — 7 уроков в день.", font_data),
    ("  • Лист «Кабинеты» (в самом конце): содержит все 24 кабинета школы с закрепленными учителями (Физика — Русина Н.В., Химия — Бабич И.Д., Математика — Пащенко Н.А., Английский — Акауова И.Н., Информатика — Чифин Д.Н., Робототехника и 4 начальных кабинета).", font_data),
    ("", font_data),
    ("После заполнения сохраните файл и загрузите его в VectorWorkspace для автоматического расчета расписания!", font_bold),
]

for row_idx, (text, f_style) in enumerate(info_texts, start=1):
    cell = ws_info.cell(row=row_idx, column=1, value=text)
    cell.font = f_style
ws_info.column_dimensions["A"].width = 110

# ═══════════════════════════════════════════════════════════════════════════
# SHEET 3: 👥 КЛАССЫ
# ═══════════════════════════════════════════════════════════════════════════
ws_classes = wb.create_sheet(title="👥 Классы")
ws_classes.views.sheetView[0].showGridLines = True

classes_headers = ["№", "Класс", "Параллель", "Литера", "Тип класса", "Смена", "Статус перехода"]
for col_idx, h in enumerate(classes_headers, start=1):
    c = ws_classes.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_navy
    c.alignment = align_center
    c.border = thin_border

for r_idx, c_info in enumerate(new_classes_info, start=2):
    ws_classes.cell(row=r_idx, column=1, value=r_idx - 1).alignment = align_center
    ws_classes.cell(row=r_idx, column=2, value=c_info["id"]).alignment = align_center
    ws_classes.cell(row=r_idx, column=3, value=c_info["grade"]).alignment = align_center
    ws_classes.cell(row=r_idx, column=4, value=c_info["letter"]).alignment = align_center
    ws_classes.cell(row=r_idx, column=5, value=c_info["type"]).alignment = align_left
    ws_classes.cell(row=r_idx, column=6, value=c_info["shift"]).alignment = align_center
    ws_classes.cell(row=r_idx, column=7, value=c_info["note"]).alignment = align_left
    for col_idx in range(1, 8):
        cell = ws_classes.cell(row=r_idx, column=col_idx)
        cell.font = font_data
        cell.border = thin_border
        if r_idx % 2 == 1:
            cell.fill = fill_zebra

for col in ws_classes.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_classes.column_dimensions[col_letter].width = max(max_len + 4, 12)

# ═══════════════════════════════════════════════════════════════════════════
# SHEET 4: 👨‍🏫 УЧИТЕЛЯ
# ═══════════════════════════════════════════════════════════════════════════
ws_teachers = wb.create_sheet(title="👨‍🏫 Учителя")
ws_teachers.views.sheetView[0].showGridLines = True

teacher_headers = ["№", "ФИО Учителя", "Основные предметы", "Закрепленный кабинет", "Макс. уроков в день", "Совместитель (Да/Нет)", "Дни и уроки приезда (свободный текст)"]
for col_idx, h in enumerate(teacher_headers, start=1):
    c = ws_teachers.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_navy
    c.alignment = align_center
    c.border = thin_border

base_room_by_tname = {}
for r in rooms_data:
    if r["teacher"] and r["teacher"] != "—":
        base_room_by_tname[r["teacher"]] = r["name"]

for r_idx, t in enumerate(teachers_list, start=2):
    tname = t["full_name"]
    broom = base_room_by_tname.get(tname, "")
    is_comb = "Да" if t.get("is_combined") else "Нет"
    
    ws_teachers.cell(row=r_idx, column=1, value=r_idx - 1).alignment = align_center
    ws_teachers.cell(row=r_idx, column=2, value=tname).alignment = align_left
    ws_teachers.cell(row=r_idx, column=3, value="").alignment = align_left
    ws_teachers.cell(row=r_idx, column=4, value=broom).alignment = align_left
    ws_teachers.cell(row=r_idx, column=5, value=7).alignment = align_center
    ws_teachers.cell(row=r_idx, column=6, value=is_comb).alignment = align_center
    ws_teachers.cell(row=r_idx, column=7, value="").alignment = align_left

    for col_idx in range(1, 8):
        cell = ws_teachers.cell(row=r_idx, column=col_idx)
        cell.font = font_data
        cell.border = thin_border
        if r_idx % 2 == 1:
            cell.fill = fill_zebra

start_empty = len(teachers_list) + 2
for r_idx in range(start_empty, start_empty + 5):
    ws_teachers.cell(row=r_idx, column=1, value=r_idx - 1).alignment = align_center
    ws_teachers.cell(row=r_idx, column=2, value="[Впишите нового учителя]").alignment = align_left
    ws_teachers.cell(row=r_idx, column=3, value="").alignment = align_left
    ws_teachers.cell(row=r_idx, column=4, value="").alignment = align_left
    ws_teachers.cell(row=r_idx, column=5, value=7).alignment = align_center
    ws_teachers.cell(row=r_idx, column=6, value="Нет").alignment = align_center
    ws_teachers.cell(row=r_idx, column=7, value="").alignment = align_left
    for col_idx in range(1, 8):
        cell = ws_teachers.cell(row=r_idx, column=col_idx)
        cell.font = font_muted
        cell.border = thin_border
        cell.fill = fill_empty_row

for col in ws_teachers.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_teachers.column_dimensions[col_letter].width = max(max_len + 4, 14)
ws_teachers.column_dimensions["B"].width = 24
ws_teachers.column_dimensions["C"].width = 30
ws_teachers.column_dimensions["D"].width = 34
ws_teachers.column_dimensions["G"].width = 36

# ═══════════════════════════════════════════════════════════════════════════
# SHEET 5: 🔗 КЛАСС-КОМПЛЕКТЫ (MULTI-CLASS SUPPORT: 2, 3 OR MORE CLASSES)
# ═══════════════════════════════════════════════════════════════════════════
ws_joint = wb.create_sheet(title="🔗 Класс-комплекты")
ws_joint.views.sheetView[0].showGridLines = True

joint_headers = [
    "№",
    "Название класс-комплекта",
    "Объединяемые классы",
    "Предмет(ы)",
    "Основной учитель",
    "Часов в нед.",
    "Тип кабинета",
    "Примечание",
]

for col_idx, h in enumerate(joint_headers, start=1):
    c = ws_joint.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_navy
    c.alignment = align_center
    c.border = thin_border

joint_rows = [
    # 3+ classes joint complexes (Multi-class example: Music for 1a, 2a, 4a)
    {"name": "1а + 2а + 4а (Музыка)", "classes": "1а, 2а, 4а", "subjs": "Музыка", "t": "Шкиря Ю.Н.", "hrs": 1, "room": "Стандартный", "note": "Объединение 3 начальных классов на Музыку"},
    {"name": "1б + 2б + 4б (Музыка)", "classes": "1б, 2б, 4б", "subjs": "Музыка", "t": "Шкиря Ю.Н.", "hrs": 1, "room": "Стандартный", "note": "Объединение 3 начальных классов на Музыку"},
    
    # LUO class-complexes (shifted +1 year)
    {"name": "8б + 8б ЛУО (Математика)", "classes": "8б, 8б ЛУО", "subjs": "Алгебра / Геометрия / Математика", "t": "Бабич И.Д.", "hrs": 5, "room": "Стандартный", "note": "Синхронизация Алгебры (3ч) и Геометрии (2ч) с Математикой (5ч)"},
    {"name": "8б + 8б ЛУО (Русский язык)", "classes": "8б, 8б ЛУО", "subjs": "Русский язык", "t": "Жолдан А.Б.", "hrs": 3, "room": "Языковой", "note": "Совмещенный урок"},
    {"name": "8б + 8б ЛУО (Литература)", "classes": "8б, 8б ЛУО", "subjs": "Русская литература / Чтение", "t": "Темерканова О.В.", "hrs": 2, "room": "Стандартный", "note": "Совмещенный урок"},
    {"name": "8б + 8б ЛУО (История Казахстана)", "classes": "8б, 8б ЛУО", "subjs": "История Казахстана", "t": "Досаева М.Х.", "hrs": 2, "room": "Стандартный", "note": "Совмещенный урок"},
    {"name": "8б + 8б ЛУО (Естествознание / Биология)", "classes": "8б, 8б ЛУО", "subjs": "Биология / Естествознание", "t": "Красноперова О.В.", "hrs": 2, "room": "Биология / Лаборатория", "note": "Совмещенный урок"},
    {"name": "8б + 8б ЛУО (Физкультура)", "classes": "8б, 8б ЛУО", "subjs": "Физическая культура", "t": "Ольков Г.В.", "hrs": 3, "room": "Спортзал", "note": "В спортзале"},
    {"name": "6б + 6 ЛУО (Русский язык и Лит.)", "classes": "6б, 6 ЛУО", "subjs": "Русский язык / Литература", "t": "Дмитриев Е.В.", "hrs": 5, "room": "Языковой", "note": "Совмещенный урок"},
    {"name": "6б + 6 ЛУО (Физкультура)", "classes": "6б, 6 ЛУО", "subjs": "Физическая культура", "t": "Ольков Г.В.", "hrs": 3, "room": "Спортзал", "note": "В спортзале"},
    {"name": "6б + 6 ЛУО (Музыка)", "classes": "6б, 6 ЛУО", "subjs": "Музыка", "t": "Шкиря Ю.Н.", "hrs": 1, "room": "Стандартный", "note": "Совмещенный урок"},
    {"name": "7 + 7 ЛУО (Биология / Естествознание)", "classes": "7, 7 ЛУО", "subjs": "Биология / Естествознание", "t": "Иващенко В.О.", "hrs": 2, "room": "Биология / Лаборатория", "note": "Совмещенный урок"},
    {"name": "7 + 7 ЛУО (Музыка)", "classes": "7, 7 ЛУО", "subjs": "Музыка", "t": "Шкиря Ю.Н.", "hrs": 1, "room": "Стандартный", "note": "Совмещенный урок"},
    {"name": "4б + 4 ЛУО (Физкультура)", "classes": "4б, 4 ЛУО", "subjs": "Физическая культура", "t": "Ольков Г.В.", "hrs": 3, "room": "Спортзал", "note": "В спортзале"},
    {"name": "4б + 4 ЛУО (Музыка)", "classes": "4б, 4 ЛУО", "subjs": "Музыка", "t": "Шкиря Ю.Н.", "hrs": 1, "room": "Стандартный", "note": "Совмещенный урок"},
    
    # Kazakh & General class-complexes templates
    {"name": "8а + 8ә (Физкультура)", "classes": "8а, 8ә", "subjs": "Физическая культура / Дене шынықтыру", "t": "Ольков Г.В.", "hrs": 3, "room": "Спортзал", "note": "Спортзал (каз. и рус. классы)"},
    {"name": "8а + 8ә (Художественный труд)", "classes": "8а, 8ә", "subjs": "Художественный труд / Көркем еңбек", "t": "Сидоров В.Б.", "hrs": 2, "room": "Мастерская", "note": "Мастерская"},
]

for r_idx, j in enumerate(joint_rows, start=2):
    ws_joint.cell(row=r_idx, column=1, value=r_idx - 1).alignment = align_center
    ws_joint.cell(row=r_idx, column=2, value=j["name"]).alignment = align_left
    ws_joint.cell(row=r_idx, column=3, value=j["classes"]).alignment = align_center
    ws_joint.cell(row=r_idx, column=4, value=j["subjs"]).alignment = align_left
    ws_joint.cell(row=r_idx, column=5, value=j["t"]).alignment = align_left
    ws_joint.cell(row=r_idx, column=6, value=j["hrs"]).alignment = align_center
    ws_joint.cell(row=r_idx, column=7, value=j["room"]).alignment = align_left
    ws_joint.cell(row=r_idx, column=8, value=j["note"]).alignment = align_left
    for col_idx in range(1, 9):
        cell = ws_joint.cell(row=r_idx, column=col_idx)
        cell.font = font_data
        cell.border = thin_border
        if r_idx % 2 == 1:
            cell.fill = fill_zebra

start_empty_j = len(joint_rows) + 2
for r_idx in range(start_empty_j, start_empty_j + 6):
    ws_joint.cell(row=r_idx, column=1, value=r_idx - 1).alignment = align_center
    ws_joint.cell(row=r_idx, column=2, value="[Новый класс-комплект]").alignment = align_left
    ws_joint.cell(row=r_idx, column=3, value="").alignment = align_center
    ws_joint.cell(row=r_idx, column=4, value="").alignment = align_left
    ws_joint.cell(row=r_idx, column=5, value="").alignment = align_left
    ws_joint.cell(row=r_idx, column=6, value="").alignment = align_center
    ws_joint.cell(row=r_idx, column=7, value="Стандартный").alignment = align_left
    ws_joint.cell(row=r_idx, column=8, value="").alignment = align_left
    for col_idx in range(1, 9):
        cell = ws_joint.cell(row=r_idx, column=col_idx)
        cell.font = font_muted
        cell.border = thin_border
        cell.fill = fill_empty_row

for col in ws_joint.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_joint.column_dimensions[col_letter].width = max(max_len + 4, 12)
ws_joint.column_dimensions["B"].width = 34
ws_joint.column_dimensions["C"].width = 24
ws_joint.column_dimensions["D"].width = 36
ws_joint.column_dimensions["E"].width = 22
ws_joint.column_dimensions["H"].width = 34

# ═══════════════════════════════════════════════════════════════════════════
# SHEET 6: 🏫 КАБИНЕТЫ (В САМОМ КОНЦЕ)
# ═══════════════════════════════════════════════════════════════════════════
ws_rooms = wb.create_sheet(title="🏫 Кабинеты")
ws_rooms.views.sheetView[0].showGridLines = True

room_headers = ["№ (Индекс)", "Идентификатор", "Название кабинета", "Тип помещения", "Закрепленный учитель"]
for col_idx, h in enumerate(room_headers, start=1):
    c = ws_rooms.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_navy
    c.alignment = align_center
    c.border = thin_border

for r_idx, r in enumerate(rooms_data, start=2):
    ws_rooms.cell(row=r_idx, column=1, value=r["index"]).alignment = align_center
    ws_rooms.cell(row=r_idx, column=2, value=r["id"]).alignment = align_center
    ws_rooms.cell(row=r_idx, column=3, value=r["name"]).alignment = align_left
    ws_rooms.cell(row=r_idx, column=4, value=r["type"]).alignment = align_left
    ws_rooms.cell(row=r_idx, column=5, value=r["teacher"]).alignment = align_left
    for col_idx in range(1, 6):
        cell = ws_rooms.cell(row=r_idx, column=col_idx)
        cell.font = font_data
        cell.border = thin_border
        if r_idx % 2 == 1:
            cell.fill = fill_zebra

for col in ws_rooms.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_rooms.column_dimensions[col_letter].width = max(max_len + 4, 14)
ws_rooms.column_dimensions["C"].width = 34
ws_rooms.column_dimensions["D"].width = 24
ws_rooms.column_dimensions["E"].width = 26

# ═══════════════════════════════════════════════════════════════════════════
# SHEET 2: 📚 УЧЕБНЫЙ ПЛАН (НАГРУЗКА) — ГЛАВНЫЙ РАБОЧИЙ ЛИСТ
# ═══════════════════════════════════════════════════════════════════════════
ws_curr = wb.create_sheet(title="📚 Учебный план (Нагрузка)", index=1)
ws_curr.views.sheetView[0].showGridLines = True

curr_headers = [
    "№",
    "Класс",
    "Предмет (по ТУП)",
    "Часов в нед.",
    "Основной учитель",
    "Деление на подгруппы?",
    "Второй учитель (при делении)",
    "Класс-комплект (Совмещение)",
    "Тип кабинета",
    "Примечание",
]

for col_idx, h in enumerate(curr_headers, start=1):
    c = ws_curr.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_navy
    c.alignment = align_center
    c.border = thin_border

old_class_by_new = {v: k for k, v in class_shift_map.items()}

old_curr_by_class = {}
for entry in curriculum:
    cid = entry["class_id"]
    old_curr_by_class.setdefault(cid, []).append(entry)

primary_tup = {
    1: [
        ("Обучение грамоте", 6, "Стандартный", "Нет"),
        ("Математика", 4, "Стандартный", "Нет"),
        ("Казахский язык", 2, "Языковой", "Да"),
        ("Естествознание", 1, "Стандартный", "Нет"),
        ("Познание мира", 1, "Стандартный", "Нет"),
        ("Музыка", 1, "Стандартный", "Нет"),
        ("Изобразительное искусство", 1, "Стандартный", "Нет"),
        ("Трудовое обучение / Труд", 1, "Мастерская", "Да"),
        ("Физическая культура", 3, "Спортзал", "Нет"),
    ],
    2: [
        ("Русский язык", 4, "Стандартный", "Нет"),
        ("Литературное чтение", 3, "Стандартный", "Нет"),
        ("Математика", 4, "Стандартный", "Нет"),
        ("Казахский язык", 2, "Языковой", "Да"),
        ("Иностранный (Английский) язык", 2, "Языковой", "Да"),
        ("Естествознание", 1, "Стандартный", "Нет"),
        ("Познание мира", 1, "Стандартный", "Нет"),
        ("Музыка", 1, "Стандартный", "Нет"),
        ("Изобразительное искусство", 1, "Стандартный", "Нет"),
        ("Трудовое обучение / Труд", 1, "Мастерская", "Да"),
        ("Физическая культура", 3, "Спортзал", "Нет"),
    ]
}

curr_row_idx = 2
item_counter = 1

joint_remap = {
    "jl_бабич_ид_7б_math": "8б + 8б ЛУО (Математика)",
    "jl_жолдан_аб_7б_rus": "8б + 8б ЛУО (Русский язык)",
    "jl_темерканова_ов_7б_lit": "8б + 8б ЛУО (Литература)",
    "jl_досаева_мх_7б_hist": "8б + 8б ЛУО (История Казахстана)",
    "jl_красноперова_ов_7б_science": "8б + 8б ЛУО (Естествознание / Биология)",
    "jl_ольков_гв_7б_pe": "8б + 8б ЛУО (Физкультура)",
    "jl_дмитриев_ев_5_rus": "6б + 6 ЛУО (Русский язык и Лит.)",
    "jl_ольков_гв_5_pe": "6б + 6 ЛУО (Физкультура)",
    "jl_шкиря_юн_5_music": "6б + 6 ЛУО (Музыка)",
    "jl_иващенко_во_6_science": "7 + 7 ЛУО (Биология / Естествознание)",
    "jl_шкиря_юн_6_music": "7 + 7 ЛУО (Музыка)",
    "jl_ольков_гв_3_pe": "4б + 4 ЛУО (Физкультура)",
    "jl_шкиря_юн_3_music": "4б + 4 ЛУО (Музыка)",
}

for cls in new_classes_info:
    new_cid = cls["id"]
    grade = cls["grade"]
    old_cid = old_class_by_new.get(new_cid)
    
    # Section Header row for Class
    ws_curr.merge_cells(start_row=curr_row_idx, start_column=1, end_row=curr_row_idx, end_column=10)
    hdr_cell = ws_curr.cell(row=curr_row_idx, column=1, value=f"Класс {new_cid}  ({cls['type']}, {cls['shift']}) — {cls['note']}")
    hdr_cell.font = font_section
    hdr_cell.fill = fill_blue
    hdr_cell.alignment = align_left
    for col_i in range(1, 11):
        ws_curr.cell(row=curr_row_idx, column=col_i).border = thin_border
    curr_row_idx += 1

    entries_to_add = []
    
    if old_cid and old_cid in old_curr_by_class:
        for entry in old_curr_by_class[old_cid]:
            sid = entry["subject_id"]
            sname = get_subj_name(sid)
            hrs = entry["hours_per_week"]
            tid = entry.get("teacher_id")
            tname = teacher_name_by_id.get(tid, tid) if tid else ""
            t2id = entry.get("split_teacher2_id")
            t2name = teacher_name_by_id.get(t2id, t2id) if t2id else ""
            
            # Primary grades 1-4: teacher empty!
            if grade <= 4:
                tname = ""
                t2name = ""
            
            old_jid = entry.get("joint_lesson_id")
            new_joint_str = joint_remap.get(old_jid, "")
            
            # Multi-class music link for 1a, 2a, 4a
            if sname == "Музыка" and new_cid in ("1а", "2а", "4а"):
                new_joint_str = "1а + 2а + 4а (Музыка)"
            elif sname == "Музыка" and new_cid in ("1б", "2б", "4б"):
                new_joint_str = "1б + 2б + 4б (Музыка)"

            req_room = get_required_room_type(sname, sid)
            is_spl = "Да" if t2name or is_split_subj(sname, sid) == "Да" else "Нет"
            
            entries_to_add.append({
                "subj": sname,
                "hrs": hrs,
                "t1": tname,
                "is_split": is_spl,
                "t2": t2name,
                "joint": new_joint_str,
                "room": req_room,
                "note": ""
            })
    elif grade in primary_tup:
        for sname, hrs, rtype, is_spl in primary_tup[grade]:
            joint_str = ""
            if sname == "Музыка" and new_cid in ("1а", "2а", "4а"):
                joint_str = "1а + 2а + 4а (Музыка)"
            elif sname == "Музыка" and new_cid in ("1б", "2б", "4б"):
                joint_str = "1б + 2б + 4б (Музыка)"

            entries_to_add.append({
                "subj": sname,
                "hrs": hrs,
                "t1": "",
                "is_split": is_spl,
                "t2": "",
                "joint": joint_str,
                "room": rtype,
                "note": ""
            })

    for item in entries_to_add:
        ws_curr.cell(row=curr_row_idx, column=1, value=item_counter).alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=2, value=new_cid).alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=3, value=item["subj"]).alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=4, value=item["hrs"]).alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=5, value=item["t1"]).alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=6, value=item["is_split"]).alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=7, value=item["t2"]).alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=8, value=item["joint"]).alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=9, value=item["room"]).alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=10, value=item["note"]).alignment = align_left

        for c_i in range(1, 11):
            cell = ws_curr.cell(row=curr_row_idx, column=c_i)
            cell.font = font_data
            cell.border = thin_border
            if curr_row_idx % 2 == 1:
                cell.fill = fill_zebra
        
        curr_row_idx += 1
        item_counter += 1

    # 3 empty rows for new subjects
    for _ in range(3):
        ws_curr.cell(row=curr_row_idx, column=1, value=item_counter).alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=2, value=new_cid).alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=3, value="[Новый предмет]").alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=4, value="").alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=5, value="").alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=6, value="Нет").alignment = align_center
        ws_curr.cell(row=curr_row_idx, column=7, value="").alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=8, value="").alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=9, value="Стандартный").alignment = align_left
        ws_curr.cell(row=curr_row_idx, column=10, value="").alignment = align_left

        for c_i in range(1, 11):
            cell = ws_curr.cell(row=curr_row_idx, column=c_i)
            cell.font = font_muted
            cell.border = thin_border
            cell.fill = fill_empty_row

        curr_row_idx += 1
        item_counter += 1

for col in ws_curr.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_curr.column_dimensions[col_letter].width = max(max_len + 3, 10)

ws_curr.column_dimensions["A"].width = 6
ws_curr.column_dimensions["B"].width = 12
ws_curr.column_dimensions["C"].width = 34
ws_curr.column_dimensions["D"].width = 14
ws_curr.column_dimensions["E"].width = 24
ws_curr.column_dimensions["F"].width = 22
ws_curr.column_dimensions["G"].width = 24
ws_curr.column_dimensions["H"].width = 34
ws_curr.column_dimensions["I"].width = 24
ws_curr.column_dimensions["J"].width = 20

# Data Validations
dv_split = DataValidation(type="list", formula1='"Да,Нет"', allow_blank=True)
ws_curr.add_data_validation(dv_split)
dv_split.add(f"F2:F{curr_row_idx}")

dv_rooms = DataValidation(type="list", formula1='"Стандартный,Спортзал,Информатика,Робототехника,Химия / Лаборатория,Физика / Лаборатория,Биология / Лаборатория,Языковой,Мастерская"', allow_blank=True)
ws_curr.add_data_validation(dv_rooms)
dv_rooms.add(f"I2:I{curr_row_idx}")

max_t_row = len(teachers_list) + 1 + 5
dv_teacher = DataValidation(type="list", formula1=f"'👨‍🏫 Учителя'!$B$2:$B${max_t_row}", allow_blank=True)
ws_curr.add_data_validation(dv_teacher)
dv_teacher.add(f"E2:E{curr_row_idx}")
dv_teacher.add(f"G2:G{curr_row_idx}")

max_j_row = len(joint_rows) + 1 + 6
dv_joint = DataValidation(type="list", formula1=f"'🔗 Класс-комплекты'!$B$2:$B${max_j_row}", allow_blank=True)
ws_curr.add_data_validation(dv_joint)
dv_joint.add(f"H2:H{curr_row_idx}")

output_path.parent.mkdir(parents=True, exist_ok=True)
wb.save(output_path)
print(f"Successfully generated clean Excel workbook at:\n{output_path}")
